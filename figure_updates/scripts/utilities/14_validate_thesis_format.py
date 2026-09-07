#!/usr/bin/env python3
"""Validate the requested typography and numbering of the final thesis DOCX."""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.text.paragraph import Paragraph


EXPECTED_FONT = "Times New Roman"
EXPECTED_FIGURES = list(range(9, 72))
FIGURE_CAPTION_RE = re.compile(r"^Figura\s+(\d+)\.", re.IGNORECASE)
FIGURE_NOTE_RE = re.compile(r"\bNota\s+de\s+las\s+Figuras\b", re.IGNORECASE)


def effective_size(paragraph, run) -> float | None:
    value = run.font.size or run.style.font.size or paragraph.style.font.size
    return round(value.pt, 2) if value is not None else None


def effective_font(document, paragraph, run) -> str | None:
    return (
        run.font.name
        or run.style.font.name
        or paragraph.style.font.name
        or document.styles["Normal"].font.name
    )


def effective_bold(paragraph, run) -> bool | None:
    return run.bold if run.bold is not None else paragraph.style.font.bold


def effective_spacing(document, paragraph) -> float | None:
    value = paragraph.paragraph_format.line_spacing
    if value is None:
        value = paragraph.style.paragraph_format.line_spacing
    if value is None:
        value = document.styles["Normal"].paragraph_format.line_spacing
    return round(float(value), 2) if isinstance(value, (int, float)) else None


def all_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from all_table_paragraphs(nested)


def element_has_image(element) -> bool:
    """Return whether a body element contains an embedded or linked image."""
    return any(
        descendant.tag.rsplit("}", 1)[-1] in {"blip", "imagedata"}
        for descendant in element.iter()
    )


def validate(docx_path: Path, pdf_path: Path | None) -> dict:
    document = Document(docx_path)
    errors: list[str] = []

    expected_styles = {
        "Normal": (12.0, None),
        "Heading 1": (14.0, True),
        "Heading 2": (13.0, True),
        "Heading 3": (13.0, True),
        "Caption": (12.0, None),
    }
    style_audit = {}
    for name, (size, bold) in expected_styles.items():
        style = document.styles[name]
        actual_size = round(style.font.size.pt, 2) if style.font.size else None
        style_audit[name] = {
            "font": style.font.name,
            "size_pt": actual_size,
            "bold": style.font.bold,
            "line_spacing": style.paragraph_format.line_spacing,
        }
        if style.font.name != EXPECTED_FONT or actual_size != size:
            errors.append(
                f"Estilo {name}: fuente/tamaño {style.font.name}/{actual_size}"
            )
        if bold is not None and style.font.bold is not bold:
            errors.append(f"Estilo {name}: negrita={style.font.bold}")

    caption_style_spacing = document.styles["Caption"].paragraph_format.line_spacing
    if (
        not isinstance(caption_style_spacing, (int, float))
        or round(float(caption_style_spacing), 2) != 1.5
    ):
        errors.append(f"Estilo Caption: interlineado={caption_style_spacing}")

    captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption" and paragraph.text.strip()
    ]
    for paragraph in captions:
        if effective_spacing(document, paragraph) != 1.5:
            errors.append(f"Pie sin interlineado 1,5: {paragraph.text[:80]}")
        for run in paragraph.runs:
            if not run.text:
                continue
            if effective_font(document, paragraph, run) != EXPECTED_FONT:
                errors.append(f"Pie sin Times New Roman: {paragraph.text[:80]}")
                break

        for run in paragraph.runs:
            if run.text and effective_size(paragraph, run) != 12.0:
                errors.append(f"Pie sin 12 pt: {paragraph.text[:80]}")
                break

    body_elements = list(document.element.body.iterchildren())
    target_figure_entries = []
    for element_index, element in enumerate(body_elements):
        if element.tag.rsplit("}", 1)[-1] != "p":
            continue
        paragraph = Paragraph(element, document._body)
        match = FIGURE_CAPTION_RE.match(paragraph.text.strip())
        if match and int(match.group(1)) in EXPECTED_FIGURES:
            target_figure_entries.append(
                (int(match.group(1)), element_index, paragraph)
            )

    figure_numbers = [number for number, _, _ in target_figure_entries]
    if figure_numbers != EXPECTED_FIGURES:
        errors.append(
            f"Figuras 9–71 ausentes, duplicadas o fuera de orden: {figure_numbers}"
        )

    for number, element_index, paragraph in target_figure_entries:
        if paragraph.style.name != "Caption":
            errors.append(f"Figura {number}: el pie no usa el estilo Caption")

        image_index = element_index - 1
        if image_index < 0 or not element_has_image(body_elements[image_index]):
            errors.append(
                f"Figura {number}: falta una imagen inmediatamente antes del pie"
            )
            continue

        narrative_index = image_index - 1
        narrative = None
        if narrative_index >= 0:
            narrative_element = body_elements[narrative_index]
            if narrative_element.tag.rsplit("}", 1)[-1] == "p":
                narrative = Paragraph(narrative_element, document._body)
        narrative_text = narrative.text.strip() if narrative is not None else ""
        if not narrative_text:
            errors.append(
                f"Figura {number}: falta un párrafo narrativo no vacío antes de la imagen"
            )
        elif FIGURE_CAPTION_RE.match(narrative_text):
            errors.append(
                f"Figura {number}: el párrafo anterior a la imagen comienza con "
                "«Figura N.»"
            )

    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading ") and paragraph.text.strip()
    ]
    for paragraph in headings:
        level = int(paragraph.style.name.rsplit(" ", 1)[1])
        expected_size = 14.0 if level == 1 else 13.0
        for run in paragraph.runs:
            if not run.text:
                continue
            if effective_size(paragraph, run) != expected_size:
                errors.append(
                    f"Encabezado con tamaño incorrecto: {paragraph.text[:80]}"
                )
                break
            if effective_bold(paragraph, run) is not True:
                errors.append(f"Encabezado sin negrita: {paragraph.text[:80]}")
                break

    reference_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if re.match(r"^\[\d+\]\s", paragraph.text)
    ]
    reference_numbers = [
        int(re.match(r"^\[(\d+)\]", paragraph.text).group(1))
        for paragraph in reference_paragraphs
    ]
    if not reference_numbers:
        errors.append("No se encontraron referencias numeradas")
    elif reference_numbers != list(range(1, len(reference_numbers) + 1)):
        errors.append(
            f"Referencias no secuenciales: {reference_numbers[:3]}...{reference_numbers[-3:]}"
        )
    for paragraph in reference_paragraphs:
        if effective_spacing(document, paragraph) != 1.5:
            errors.append(f"Referencia sin interlineado 1,5: {paragraph.text[:60]}")
        if any(
            run.text and effective_size(paragraph, run) != 12.0
            for run in paragraph.runs
        ):
            errors.append(f"Referencia sin 12 pt: {paragraph.text[:60]}")

    main_text = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "VII. REFERENCIAS":
            break
        main_text.append(paragraph.text)
    joined_main = "\n".join(main_text)
    residual_patterns = re.findall(
        r"(?:[A-ZÁÉÍÓÚÜÑ][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+ et al\.\s*\(20\d{2}[a-z]?\)"
        r"|\((?:World Health Organization|U\.S\. Food and Drug Administration|"
        r"[A-ZÁÉÍÓÚÜÑ][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+ et al\.),\s*20\d{2})",
        joined_main,
    )
    if residual_patterns:
        errors.append(f"Citas autor-año residuales: {residual_patterns[:5]}")

    table_paragraphs = [
        paragraph
        for table in document.tables
        for paragraph in all_table_paragraphs(table)
    ]
    figure_note_paragraphs = [
        paragraph.text[:80]
        for paragraph in [*document.paragraphs, *table_paragraphs]
        if FIGURE_NOTE_RE.search(paragraph.text)
    ]
    if figure_note_paragraphs:
        errors.append(
            "Se encontraron párrafos «Nota de las Figuras»: "
            f"{figure_note_paragraphs}"
        )

    bad_table_runs = [
        paragraph.text[:60]
        for paragraph in table_paragraphs
        if any(
            run.text and effective_size(paragraph, run) != 12.0
            for run in paragraph.runs
        )
    ]
    if bad_table_runs:
        errors.append(f"Celdas sin 12 pt: {bad_table_runs[:5]}")

    bad_style_fonts = []
    for style in document.styles:
        if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        if style.font.name not in (None, EXPECTED_FONT):
            bad_style_fonts.append((style.name, style.font.name))
    if bad_style_fonts:
        errors.append(f"Fuentes de estilo ajenas: {bad_style_fonts[:8]}")

    with zipfile.ZipFile(docx_path) as archive:
        bad_member = archive.testzip()
        settings = archive.read("word/settings.xml")
        footer_parts = [
            name
            for name in archive.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        ]
        footer_xml = b"\n".join(archive.read(name) for name in footer_parts)
    if bad_member:
        errors.append(f"Miembro DOCX corrupto: {bad_member}")
    if b"updateFields" not in settings:
        errors.append("Falta updateFields en settings.xml")
    if b" PAGE " not in footer_xml:
        errors.append("Falta campo PAGE en el pie de página")

    pdf_pages = None
    if pdf_path is not None:
        if not pdf_path.is_file():
            errors.append(f"No existe el PDF: {pdf_path}")
        else:
            result = subprocess.run(
                ["pdfinfo", str(pdf_path)],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
            )
            match = re.search(r"^Pages:\s+(\d+)$", result.stdout, re.MULTILINE)
            pdf_pages = int(match.group(1)) if match else None
            if result.returncode or not pdf_pages:
                errors.append(
                    "No se pudo determinar el número de páginas del PDF: "
                    + result.stderr.strip()
                )

    return {
        "status": "PASS" if not errors else "FAIL",
        "docx": str(docx_path),
        "pdf": str(pdf_path) if pdf_path else None,
        "pdf_pages": pdf_pages,
        "paragraphs": len(document.paragraphs),
        "inline_shapes": len(document.inline_shapes),
        "captions": len(captions),
        "figures_9_71": figure_numbers,
        "figure_notes": len(figure_note_paragraphs),
        "headings": len(headings),
        "references": len(reference_paragraphs),
        "style_audit": style_audit,
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    args = parser.parse_args()
    result = validate(args.docx.resolve(), args.pdf.resolve() if args.pdf else None)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
