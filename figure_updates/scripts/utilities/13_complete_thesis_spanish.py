#!/usr/bin/env python3
"""Complete the Spanish thesis from its preserved Word source.

The build deliberately keeps the original cover/defence pages, the thesis-only
schemes, the two-bottle design and the genotyping gels.  Result figures that
have canonical Spanish paper panels are replaced with those panels.  No raw
microscopy or behavioural data are read or modified by this script.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT_SOURCES = ROOT / "manuscript_sources"
if not MANUSCRIPT_SOURCES.is_dir():
    MANUSCRIPT_SOURCES = ROOT / "thesis_and_manuscript/insumos/manuscript_sources"
if not MANUSCRIPT_SOURCES.is_dir():
    MANUSCRIPT_SOURCES = ROOT / "revision_profesional_20260906/insumos/manuscript_sources"
SOURCE = MANUSCRIPT_SOURCES / "TESIS_FINAL_NS4_source_20260830.docx"
DEFAULT_OUTPUT = ROOT / "TESIS_FINAL_NS(4).docx"
WORK = Path("/tmp/thesis_ns4_work_20260830")
MEDIA = WORK / "media" / "media"
PRESERVED_MEDIA_NAMES = ("image27.png", "image28.png")
LITERATURE = ROOT / "thesis_and_manuscript" / "insumos" / "literature_webscrap_30_08_2026"
if not LITERATURE.is_dir():
    LITERATURE = ROOT / "revision_profesional_20260906" / "insumos" / "literature_webscrap_30_08_2026"
if not LITERATURE.is_dir():
    LITERATURE = ROOT / "literature_webscrap_30_08_2026"
BUILD_DATE = "30 de agosto de 2026"
PAGE_MAP_PATH = (
    MANUSCRIPT_SOURCES /
    "TESIS_FINAL_NS4_page_map_20260830.json"
)
PAPER_TITLE = (
    "Oral allulose after prior familiarization is associated with a "
    "distinct spatial c-FOS/NPY profile without detectable POMC or glial "
    "activation in mice"
)
PAPER_AUTHORS = [
    "Nancy Segura",
    "Vinka Azócar",
    "Claudia Aguilera",
    "Constanza Sanhueza",
    "Bastían Aravena",
    "Rocío Casas",
    "Rocío Magdalena",
    "Isidora Troncoso",
    "Aracelly Quiroz",
    "Daniela Mennickent Barros",
    "Ricardo Flores",
    "Francisca Espinoza",
    "Antonia Recabal-Beyer",
    "Roberto Elizondo",
    "Mauricio D Dorfmann",
    "María de los Ángeles García",
    "Valentina González-Pecchi",
    "Carlos Farkas*",
]

BODY_FONT_SIZE = 12.0
TITLE_FONT_SIZE = 14.0
SUBTITLE_FONT_SIZE = 13.0
THESIS_LINE_SPACING = 1.5
CITATIONS = None


def render_citations(text: str) -> str:
    """Render author-year citations with the active numbered bibliography."""
    if CITATIONS is None:
        return text
    return CITATIONS.convert(text)


def set_run_font(run, name: str = "Times New Roman",
                 size: Pt | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = size
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), name)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, size=Pt(BODY_FONT_SIZE))


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_text(paragraph: Paragraph, text: str, *, bold: bool = False,
             italic: bool = False, align=None,
             size: float = BODY_FONT_SIZE) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(render_citations(text))
    run.bold = bold
    run.italic = italic
    set_run_font(run, size=Pt(size))
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    if align is not None:
        paragraph.alignment = align


def insert_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_field(paragraph: Paragraph, instruction: str) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Actualice este campo en Word o LibreOffice."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_update_fields_setting(document: Document) -> None:
    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def delete_from(document: Document, paragraph_index: int) -> None:
    start = document.paragraphs[paragraph_index]._element
    body = document._element.body
    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if deleting and child.tag != qn("w:sectPr"):
            body.remove(child)


def add_heading(document: Document, text: str, level: int = 1) -> Paragraph:
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    return paragraph


def apply_heading_style(paragraph: Paragraph, level: int) -> None:
    """Apply heading hierarchy even when the source run has direct formatting."""
    paragraph.style = paragraph.part.document.styles[f"Heading {level}"]
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    size = TITLE_FONT_SIZE if level == 1 else SUBTITLE_FONT_SIZE
    for run in paragraph.runs:
        run.bold = True
        set_run_font(run, size=Pt(size))


def add_body(document: Document, text: str, *, italic: bool = False,
             bold: bool = False) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.35)
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(render_citations(text))
    run.italic = italic
    run.bold = bold
    set_run_font(run, size=Pt(BODY_FONT_SIZE))
    return paragraph


def add_caption(document: Document, text: str) -> Paragraph:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(render_citations(text))
    set_run_font(run, size=Pt(BODY_FONT_SIZE))
    return paragraph


def add_panel_caption(document: Document, label: str, text: str) -> Paragraph:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.keep_together = True
    prefix = paragraph.add_run(label.rstrip() + " ")
    prefix.bold = True
    set_run_font(prefix, size=Pt(BODY_FONT_SIZE))
    body = paragraph.add_run(render_citations(
        re.sub(r"\s+", " ", text).strip()))
    set_run_font(body, size=Pt(BODY_FONT_SIZE))
    return paragraph


def picture_dimensions(path: Path, max_width: float = 6.05,
                       max_height: float = 7.65) -> tuple[float, float]:
    with Image.open(path) as im:
        width_px, height_px = im.size
    ratio = width_px / max(height_px, 1)
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def add_picture(document: Document, path: Path, *, label: str | None = None,
                page_break: bool = False, max_height: float = 7.65) -> Paragraph:
    if not path.is_file():
        raise FileNotFoundError(path)
    if label:
        label_p = document.add_paragraph()
        label_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        label_p.paragraph_format.space_after = Pt(3)
        label_p.paragraph_format.keep_with_next = True
        run = label_p.add_run(label)
        run.bold = True
        set_run_font(run, size=Pt(BODY_FONT_SIZE))
    width, height = picture_dimensions(path, max_height=max_height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
    if page_break:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def add_panel_picture(document: Document, path: Path, caption_label: str,
                      caption_text: str, *, max_height: float = 5.85) -> None:
    width, height = picture_dimensions(
        path, max_width=6.05, max_height=max_height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width), height=Inches(height))
    add_panel_caption(document, caption_label, caption_text)


def add_legend(document: Document, path: Path, prefix: str | None = None) -> None:
    text = path.read_text(encoding="utf-8").strip()
    blocks = [re.sub(r"\s+", " ", block.strip())
              for block in re.split(r"\n\s*\n", text) if block.strip()]
    if prefix and blocks:
        blocks[0] = re.sub(r"^(FIGURA|Figura)( suplementaria)?\s+[^.]+\.",
                           prefix, blocks[0], count=1)
    for block in blocks:
        add_caption(document, block)


def add_figure_set(document: Document, title: str, panels: list[Path],
                   legend: Path, prefix: str) -> None:
    add_heading(document, title, 3)
    for index, panel in enumerate(panels):
        letter = chr(ord("A") + index)
        add_picture(document, panel, label=f"{prefix}{letter}", page_break=True)
    add_legend(document, legend, prefix=f"{prefix}. ")
    document.add_page_break()


def style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_FONT_SIZE)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = THESIS_LINE_SPACING
    normal.paragraph_format.space_after = Pt(6)
    for level, size in (
            (1, TITLE_FONT_SIZE),
            (2, SUBTITLE_FONT_SIZE),
            (3, SUBTITLE_FONT_SIZE)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = THESIS_LINE_SPACING
        style.paragraph_format.keep_with_next = True
    for section in document.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(1.18)
        section.header_distance = Inches(0.4)
        section.footer_distance = Inches(0.45)
        footer = section.footer
        footer.is_linked_to_previous = False
        footer_element = footer._element
        # The source stores its old right-aligned page number inside a w:sdt,
        # which python-docx does not expose through footer.paragraphs.  Purge
        # the complete footer part so the old and rebuilt fields cannot stack.
        for child in list(footer_element):
            footer_element.remove(child)
        paragraph_element = OxmlElement("w:p")
        footer_element.append(paragraph_element)
        add_page_number(Paragraph(paragraph_element, footer))

    try:
        caption = document.styles["Caption"]
    except KeyError:
        caption = document.styles.add_style(
            "Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(BODY_FONT_SIZE)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    caption.paragraph_format.line_spacing = THESIS_LINE_SPACING
    caption.paragraph_format.space_after = Pt(8)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def sanitize_fonts(document: Document) -> None:
    """Force one Word-text font without altering figure raster content."""
    for style in document.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = "Times New Roman"
            rpr = style._element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn(f"w:{attribute}"), "Times New Roman")

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(iter_table_paragraphs(table))
    for section in document.sections:
        for part in (section.header, section.footer):
            paragraphs.extend(part.paragraphs)
            for table in part.tables:
                paragraphs.extend(iter_table_paragraphs(table))
    for paragraph in paragraphs:
        for run in paragraph.runs:
            set_run_font(run)

    # Some source text is stored inside structured document tags and is not
    # exposed through document.paragraphs.  Normalize those run declarations
    # directly so the defence pages and text boxes do not retain Calibri.
    roots = [document._element]
    for section in document.sections:
        roots.extend((section.header._element, section.footer._element))
    for root in roots:
        for rfonts in root.xpath(".//w:rFonts"):
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn(f"w:{attribute}"), "Times New Roman")


INTRO_REPLACEMENTS = {
    212: "La obesidad, la diabetes mellitus tipo 2 y la enfermedad hepática esteatósica asociada a disfunción metabólica constituyen problemas sanitarios de alta prevalencia y etiología multifactorial. Su desarrollo se relaciona con factores biológicos, ambientales y sociales, entre ellos la disponibilidad de alimentos de alta densidad energética, el patrón dietario y la actividad física. La Organización Mundial de la Salud recomienda reducir los azúcares libres a menos del 10 % de la energía diaria y señala un beneficio adicional probable por debajo del 5 %, aunque esta recomendación no atribuye el riesgo metabólico a un único nutriente (World Health Organization, 2015, 2024, 2025; Hall et al., 2019; Popkin y Ng, 2022).",
    213: "La conducta alimentaria surge de la interacción entre mecanismos homeostáticos, procesos de recompensa, aprendizaje, disponibilidad de alimentos y contexto social. El hambre favorece la búsqueda y el consumo de alimento, mientras que la saciación contribuye a finalizar una comida y la saciedad limita el inicio de una nueva. Estas respuestas no dependen de un circuito aislado, sino de redes distribuidas que integran señales internas y estímulos sensoriales (Campos et al., 2022; Capucho y Conde, 2022).",
    214: "Sistema homeostático y señales periféricas",
    215: "El control homeostático de la ingesta integra información sobre reservas energéticas, nutrientes circulantes y señales gastrointestinales. Leptina e insulina informan sobre el estado energético de mediano plazo, mientras que grelina, colecistoquinina, GLP-1 y PYY participan en respuestas asociadas con el ayuno y la comida. Estas señales alcanzan el encéfalo por la circulación, por aferencias vagales y por circuitos del tronco encefálico, y convergen con información sensorial y cognitiva (Jais y Brüning, 2022; Wouters d’Oplinter et al., 2022).",
    216: "El hipotálamo como centro integrador del balance energético",
    217: "El hipotálamo se ubica en la base del diencéfalo, alrededor del tercer ventrículo, y coordina funciones endocrinas, autonómicas y conductuales. En la regulación energética participan, entre otras regiones, el núcleo arqueado (ARC), el núcleo paraventricular, el núcleo ventromedial (VMN), el núcleo dorsomedial y el área hipotalámica lateral. Estas regiones reciben señales periféricas y se conectan con el tronco encefálico y circuitos de recompensa; por ello, el hipotálamo es un nodo integrador y no un regulador único de la conducta alimentaria (Campos et al., 2022; Jais y Brüning, 2022).",
    219: "Figura 1. Esquema de núcleos hipotalámicos. Corte coronal murino con la ubicación aproximada del ARC, VMN, DMH, PVN, LH, tercer ventrículo (3V) y eminencia media (EM). Esquema propio conservado de la versión original de la tesis.",
    221: "Diversidad celular del hipotálamo",
    222: "El hipotálamo contiene neuronas excitatorias e inhibitorias, astrocitos, oligodendrocitos, precursores oligodendrogliales, microglía, células vasculares, ependimarias y tanicitos. Los atlas de célula única, incluido HypoMap, muestran que estas categorías comprenden numerosos estados moleculares y subtipos. Las proporciones dependen de la región, el método de aislamiento y el criterio de anotación, por lo que no deben interpretarse como una composición universal de todo el hipotálamo (Steuernagel et al., 2022).",
    231: "Figura 2. Mapa de referencia unificado del hipotálamo murino.",
    232: "Integración de conjuntos de secuenciación de célula única con los principales tipos celulares y clústeres neuronales. Las neuronas inhibitorias se identifican principalmente por Slc32a1 y las excitatorias por Slc17a6. Adaptado de Steuernagel et al. (2022); se conserva la figura de revisión incluida en la tesis original.",
    235: "Núcleo arqueado y sistema melanocortínico",
    236: "El ARC se encuentra junto al tercer ventrículo y la eminencia media, donde la interfaz vascular y los tanicitos facilitan el intercambio de señales metabólicas. La accesibilidad no es uniforme y no equivale a una ausencia completa de barrera hematoencefálica. Esta organización permite que poblaciones del ARC respondan a hormonas y nutrientes y modulen otros núcleos hipotalámicos (Clayton et al., 2022; Jais y Brüning, 2022).",
    237: "Las neuronas que expresan proopiomelanocortina (POMC) producen péptidos derivados de POMC, incluida α-MSH. La activación de receptores melanocortínicos, en especial MC4R en neuronas postsinápticas, suele reducir la ingesta y favorecer el gasto energético. Sin embargo, las neuronas POMC son heterogéneas y su respuesta depende del estado energético, la proyección y el estímulo (Jais y Brüning, 2022; Sousa et al., 2023).",
    238: "Las neuronas que coexpresan NPY y AgRP participan de forma central en respuestas de ayuno y búsqueda de alimento. AgRP antagoniza la señal melanocortínica, mientras que NPY y GABA actúan sobre receptores y circuitos postsinápticos. Su actividad puede aumentar con grelina o baja disponibilidad de glucosa y disminuir con leptina, insulina y señales posprandiales. La expresión de c-FOS en una neurona NPY es un indicador indirecto y temporal de respuesta celular; no equivale por sí sola a una medición de hambre o ingesta.",
    240: "Figura 3. Regulación neuroendocrina de la ingesta y la homeostasis energética.",
    241: "Esquema de la integración de señales periféricas por neuronas AgRP/NPY y POMC y su proyección hacia circuitos melanocortínicos. Los signos representan efectos predominantes y no excluyen heterogeneidad celular ni conexiones adicionales. Esquema propio conservado de la versión original de la tesis (Lavoie et al., 2023; Sweeney et al., 2023).",
    242: "Células gliales y regulación metabólica",
    243: "Las células gliales mantienen la homeostasis tisular y participan en la comunicación metabólica. Astrocitos, microglía, oligodendrocitos y células ependimarias responden a cambios nutricionales y pueden modificar la disponibilidad de metabolitos, la función sináptica y el entorno inflamatorio. La evidencia disponible es principalmente preclínica y depende de la región y del modelo experimental (Frintrop et al., 2021; Jais y Brüning, 2022).",
    244: "Los astrocitos expresan receptores para señales metabólicas y regulan la captación de nutrientes, el soporte energético y el entorno sináptico. Manipulaciones quimiogenéticas y estudios de remodelación perisomática sugieren que pueden modificar la actividad de neuronas AgRP y POMC; no obstante, la dirección del efecto varía con la herramienta, la dosis, el estado nutricional y el tiempo de observación (Yang et al., 2015; Nuzzaci et al., 2020). GFAP informa sobre una parte del citoesqueleto astrocítico y no mide por sí sola toda la función o la denominada reactividad.",
    245: "La microglía es la población inmunitaria residente del sistema nervioso central y presenta un continuo de estados transcripcionales y morfológicos. La clasificación binaria entre microglía 'en reposo' y 'activada' es insuficiente. Iba1 permite visualizar células microgliales, pero su intensidad o morfología aislada no demuestra una función inflamatoria específica. Por esta razón, en este trabajo las categorías morfométricas predichas por computadora se consideran exploratorias y no equivalen a estados biológicos validados (Clayton et al., 2022; Frintrop et al., 2021).",
    246: "Consumo de azúcares y regulación metabólica",
    247: "El consumo elevado y sostenido de azúcares libres puede aumentar la densidad energética de la dieta y contribuir a alteraciones cardiometabólicas. En modelos animales, dietas ricas en azúcar también se han asociado con cambios neuronales y gliales; la magnitud y dirección dependen de la dosis, la duración, el balance energético y la especie. Estos resultados no permiten inferir que una exposición aguda produzca neuroinflamación (Teysseire et al., 2022; Patkar et al., 2021).",
    248: "La recomendación de la OMS sobre azúcares libres busca reducir caries y aumento no saludable de peso y se expresa como porcentaje de la energía total. No establece un límite universal de 50 g para todas las personas, porque el valor absoluto depende del requerimiento energético. En esta tesis, sacarosa se utiliza como comparador calórico dulce y agua como control sin azúcar (World Health Organization, 2015).",
    249: "Sustitutos del azúcar y alcance de las recomendaciones",
    250: "Los sustitutos del azúcar incluyen compuestos químicamente distintos: edulcorantes intensos, polioles y azúcares de baja disponibilidad energética. Su dulzor, absorción, metabolismo y efectos gastrointestinales no son intercambiables. En Chile, la reformulación posterior a la ley de etiquetado aumentó la presencia de edulcorantes en varias categorías de alimentos, pero la exposición dietaria y el beneficio clínico deben evaluarse por compuesto y patrón de consumo (Zancheta Ricardo et al., 2021).",
    251: "En 2023 la OMS emitió una recomendación condicional contra el uso de edulcorantes sin azúcar para controlar el peso a largo plazo. Esa guía excluye explícitamente los azúcares y derivados de azúcares de bajo aporte energético, entre ellos la alulosa; por tanto, no debe aplicarse directamente a este compuesto. Del mismo modo, resultados recientes con sucralosa sobre apetito o respuesta cerebral no pueden generalizarse a la alulosa porque difieren en estructura, dosis y destino metabólico (World Health Organization, 2023; Chakravartti et al., 2025).",
    252: "Los denominados azúcares raros son monosacáridos presentes en bajas cantidades en la naturaleza y producidos industrialmente mediante conversión enzimática. La alulosa es el miembro con mayor desarrollo alimentario, pero su clasificación como azúcar raro no demuestra eficacia terapéutica.",
    253: "Alulosa",
    254: "La D-alulosa, antes denominada D-psicosa, es el epímero en C-3 de la D-fructosa. Tiene aproximadamente el 70 % del dulzor de la sacarosa y aporta poca energía metabolizable. Para rotulado en Estados Unidos, la FDA permite usar 0,4 kcal/g y excluirla de 'azúcares totales' y 'azúcares añadidos', manteniéndola en carbohidratos totales. Esta decisión regulatoria no constituye por sí misma evidencia de beneficio clínico (U.S. Food and Drug Administration, 2020).",
    256: "Figura 4. Estructura química de D-fructosa y D-alulosa.",
    257: "Ambos monosacáridos comparten fórmula molecular y difieren en la orientación del grupo hidroxilo del carbono 3. Se conserva el esquema químico de la versión original de la tesis.",
    258: "Tras la ingesta, una fracción importante de la alulosa se absorbe en el intestino delgado y se elimina por orina con metabolismo limitado; otra fracción alcanza el colon. Dosis altas pueden producir síntomas gastrointestinales, por lo que la tolerabilidad depende de la dosis absoluta y por kilogramo de peso corporal. La exposición experimental debe describirse con concentración, volumen, vía y estado de ayuno.",
    259: "Efectos metabólicos de la alulosa",
    260: "Ensayos agudos en humanos indican que la alulosa administrada con carbohidratos puede atenuar la respuesta posprandial de glucosa y, en algunos estudios, de insulina. Una revisión sistemática y metaanálisis reciente encontró evidencia más consistente para variables posprandiales que para glucosa en ayunas, lípidos o composición corporal, donde no se demostró un beneficio robusto. La heterogeneidad de dosis, formulaciones y poblaciones limita la extrapolación (Franchi et al., 2021; Ayesh et al., 2024; Osborn et al., 2026).",
    261: "En roedores, la alulosa ha reducido ganancia de peso y alteraciones metabólicas en algunos modelos de dieta alta en grasa o diabetes. Estos trabajos aportan plausibilidad mecanística, pero no demuestran eficacia clínica ni permiten asumir que una exposición aguda en animales normopeso tendrá el mismo efecto (Chen et al., 2019; Bae et al., 2023).",
    262: "Alulosa y regulación del apetito",
    263: "Iwasaki y colaboradores administraron alulosa por vía peroral mediante sonda hacia el estómago de ratones en ayuno y observaron reducción transitoria de la ingesta con 1 y 3 g/kg, liberación de GLP-1 y participación de aferencias vagales. El efecto dependió de la vía y no se reprodujo con administración intraperitoneal. Estos datos apoyan un eje intestino-cerebro, pero proceden de un protocolo distinto del utilizado en esta tesis (Iwasaki et al., 2018).",
    264: "Alulosa y acción sobre el núcleo arqueado",
    265: "El grupo de Toshihiko Yada estudió dos contextos centrales. Yermek et al. (2022a) aplicaron alulosa directamente a neuronas aisladas del ARC y usaron inyección intracerebroventricular para el ensayo de alimentación; informaron activación de una fracción de neuronas POMC y reducción de la ingesta. En un segundo trabajo, Yermek et al. (2022b) observaron inhibición de una fracción de neuronas NPY, sensibles a grelina o a baja glucosa, y supresión de la ingesta tras inyección intracerebroventricular. Esta vía es intracerebroventricular, no intratecal, y evita la fase oral, la absorción gastrointestinal y gran parte de la señalización periférica.",
    266: "En el presente trabajo, los animales recibieron durante la adaptación la solución asignada y, tras el reemplazo por agua y un ayuno de 4 horas, recibieron 200 µL por pipeta en la cavidad oral. El contraste entre una exposición oral después de aprendizaje previo y una exposición central o directa puede modificar la señal anticipatoria, el tránsito gastrointestinal, las hormonas intestinales y el tiempo de llegada de señales al ARC. Por ello, una dirección distinta en c-FOS/NPY sería biológicamente plausible, pero no puede atribuirse a la vía sin un experimento que la compare de manera directa.",
    267: "Este estudio evaluó si la alulosa, comparada con agua y sacarosa, se asociaba con cambios conductuales y con marcadores celulares en ARC, EM y VMN. Se priorizó al animal o la jaula como unidad experimental, se conservaron los controles de calidad y se interpretaron como exploratorios los análisis con muestras pequeñas, múltiples variables o clasificadores sin validación experta.",
}


METHOD_REPLACEMENTS = {
    269: "HIPÓTESIS Y OBJETIVOS",
    270: "Hipótesis",
    271: "La exposición aguda a alulosa, después de un periodo de adaptación a la solución, modifica la respuesta de subpoblaciones celulares del hipotálamo mediobasal mediante señales periféricas y centrales, sin requerir un aumento general de marcadores gliales.",
    272: "Objetivo general",
    273: "Evaluar el efecto de la exposición a alulosa sobre la conducta de consumo y sobre marcadores neuronales y gliales del hipotálamo en modelos murinos.",
    274: "Objetivos específicos",
    275: "Comparar la respuesta c-FOS global y la respuesta de células NPY y POMC después de una exposición oral aguda a agua, sacarosa o alulosa.",
    276: "Comparar la señal GFAP, la señal Iba1 y descriptores morfométricos microgliales entre tratamientos, con delimitación anatómica aceptada mediante HIL.",
    277: "Evaluar el consumo de las soluciones desde el día 1, el consumo de alimento y el cambio de peso corporal usando la jaula como unidad experimental cuando la exposición fue compartida.",
    293: "METODOLOGÍA",
    294: "Animales y consideraciones éticas",
    295: "Se utilizaron ratones reporteros NPY-GFP (B6.FVB-Tg(Npy-hrGFP)1Lowl/J; JAX 006417) y POMC-GFP (C57BL/6J-Tg(Pomc-EGFP)1Low/J; JAX 009593), además de ratones C57BL/6J. Los animales se mantuvieron con ciclo luz/oscuridad de 12 h, alimento y agua ad libitum salvo durante los ayunos indicados, y temperatura aproximada de 25 °C en el Centro Regional de Estudios Avanzados para la Vida. Los procedimientos fueron revisados y aprobados por el comité ético-científico institucional. El archivo fuente no consignaba el número de protocolo; por ello no se inventó uno en esta versión.",
    296: "Experimento de una botella",
    297: "Se incluyeron ratones de las líneas POMC-GFP y NPY-GFP asignados a agua, sacarosa al 10 % o alulosa al 13 %. El análisis conductual final dispuso de 12 jaulas: Agua n=3, Sacarosa n=5 y Alulosa n=4; los puntos de ratón se conservaron como descripción, pero la inferencia se realizó por jaula. La cantidad de animales o secciones varió entre análisis histológicos según disponibilidad y control de calidad, y se informa en cada figura.",
    298: "Días 1–7: adaptación",
    299: "Cada grupo recibió dieta estándar y la solución asignada ad libitum durante siete días. Se registraron el peso corporal y el consumo desde el día 1. Cuando dos animales compartieron una jaula, el consumo de la botella y del alimento se consideró una medición de jaula y no se dividió en observaciones animales independientes.",
    300: "Días 8–9: reemplazo por agua",
    301: "Las soluciones dulces se reemplazaron por agua durante dos días. Se mantuvieron la dieta estándar y el registro de peso. Esta etapa separó la exposición de adaptación de la exposición aguda y redujo la presencia inmediata de solución dulce antes del desafío final.",
    302: "Día 10: exposición oral aguda",
    303: "Después de un ayuno de 4 h, cada animal recibió 200 µL de agua, sacarosa al 10 % o alulosa al 13 % mediante pipeta dirigida a la cavidad oral. En un animal de 25 g, 200 µL de alulosa al 13 % equivalen aproximadamente a 1,04 g/kg; el valor exacto depende del peso individual. Esta administración oral se diferencia de la inyección intracerebroventricular y de la superfusión de neuronas aisladas utilizadas por Yermek et al. (2022a, 2022b).",
    304: "Perfusión, fijación y obtención de cortes",
    305: "Entre 30 y 40 min después de la exposición, según la cohorte indicada en cada figura, los animales fueron anestesiados con ketamina (150 mg/kg) y xilacina (10 mg/kg) y perfundidos por vía transcardíaca con PBS seguido de paraformaldehído al 4 %. Los cerebros se posfijaron, crioprotegieron en sacarosa al 30 %, incluyeron en OCT y se cortaron en secciones coronales. Para el análisis montado se obtuvieron secciones de 20 µm y para inmunohistoquímica flotante, secciones de 30 µm. Se analizaron niveles del hipotálamo mediobasal compatibles con las coordenadas aproximadas bregma −1,23 a −2,15 mm.",
    306: "Experimento de dos botellas",
    307: "Se utilizaron 16 ratones C57BL/6J, dos por jaula, con una botella de agua y una segunda botella asignada a agua, sacarosa al 10 % o alulosa al 13 %. La cohorte final comprendió ocho jaulas: Agua n=2, Alulosa n=3 y Sacarosa n=3. Se analizaron siete días posteriores a la inicialización. En controles Agua, la fracción de la segunda botella representa asignación lateral y no preferencia por un edulcorante.",
    308: "Registro conductual y de peso",
    309: "El peso corporal se registró diariamente con una balanza de laboratorio cuya marca y modelo no estaban documentados en el archivo fuente. El consumo se calculó a partir del cambio de masa de la botella bajo la aproximación de 1 g/mL y se informó por jaula, no como ingesta individual. El alimento se pesó por jaula. Se analizaron trayectorias y valores finales por jaula para evitar pseudorreplicación.",
    310: "Genotipificación",
    311: "Se obtuvo ADN genómico a partir de aproximadamente 2 mm de cola. Cada muestra se incubó con 75 µL de solución de lisis alcalina a 95 °C durante 1 h, se enfrió y se neutralizó con 75 µL de Tris-HCl 40 mM. Tras centrifugación a 4.000 rpm durante 3 min, se recuperó el sobrenadante.",
    312: "Para NPY-GFP se usaron los partidores común 5′-TATGTGGACGGGGCAGAAGATCCAGG-3′, silvestre 5′-CCCAGCTCACATATTTATCTAGAG-3′ y transgénico 5′-GGTGCGGTTGCCGTACTGGA-3′. Para POMC-GFP se usaron 5′-CTGTCCTCAGAAAGCCTTGG-3′ y 5′-CTGAACTTGTGGCCGTTTAC-3′, junto con el control interno 5′-CAAATGTTGCTTGTCTGGTG-3′ y 5′-GTCAGTCGAGTGCACAGTTT-3′.",
    313: "La PCR se realizó con 3 µL de ADN: 95 °C por 2 min; 35 ciclos de 95 °C por 15 s, 57 °C por 15 s y 72 °C por 30 s; y extensión final a 72 °C por 5 min. Los tamaños esperados fueron 400 pb para NPY-GFP y 500 pb para su control, y 173 pb para POMC-GFP y 324 pb para su control.",
    314: "Electroforesis en gel de agarosa",
    315: "Los productos se separaron en geles de agarosa al 2 % en TAE 1X a 100 V durante 40 min. Se utilizó colorante MaestroSafe, tampón de carga 6X y un marcador de 100 pb–4 kb. Los geles se visualizaron en un transiluminador Omega Lum G.",
    316: "Inmunohistoquímica fluorescente montada",
    317: "c-FOS se utilizó como marcador de respuesta celular inmediata. Su expresión depende del estímulo y del tiempo y no constituye una lectura directa de disparo neuronal, hambre o función secretora. En secciones montadas se evaluaron c-FOS, NeuN, DAPI y la fluorescencia reportera NPY-GFP; la falta de señal POMC-GFP utilizable impidió emplear esa línea para la variable POMC en esta cohorte.",
    318: "Las secciones se lavaron en tampón Tris-fosfato, se bloquearon durante 1 h con Triton X-100 al 0,3 % y suero normal de cabra al 10 %, y se incubaron durante la noche a 4 °C con anti-c-FOS de conejo (1:100; Cell Signaling 31254) y anti-NeuN de ratón (1:200–1:250; Sigma MAB377).",
    319: "Después de los lavados se incubaron anticuerpos secundarios apropiados, incluido anti-conejo Alexa Fluor 647 (1:200; Jackson 711-605-152) y anti-ratón Alexa Fluor 488 (1:200; Cell Signaling 4408), durante 2 h en oscuridad. Los núcleos se marcaron con DAPI y las secciones se montaron con Fluoromount-G.",
    321: "Figura 7. Diseño experimental de una botella.",
    322: "Panel canónico en español del manuscrito. Después de siete días de adaptación y dos días de reemplazo por agua, los animales ayunaron 4 h y recibieron la solución por vía oral mediante pipeta. La eutanasia y fijación se realizaron 30–40 min después. Este esquema reemplaza la versión antigua del documento y conserva la secuencia utilizada por el análisis final.",
    323: "Inmunohistoquímica fluorescente flotante",
    324: "Las secciones flotantes de 30 µm se lavaron en PBS, se bloquearon con suero normal de burro al 2,5 % y Triton X-100 al 0,5 %, y se incubaron durante la noche a 4 °C con anti-c-FOS de cabra (1:1.000; Synaptic Systems 226308), anti-POMC de cobayo (1:4.000; Phoenix G-029-30), anti-Iba1 de conejo (1:2.000; Wako 019-19741) y anti-GFAP-Cy3 de ratón (1:1.000–1:2.000; Sigma C9205), según el ensayo. Se utilizaron anticuerpos secundarios compatibles y medio de montaje con DAPI.",
    326: "Figura 8. Diseño experimental de dos botellas.",
    327: "Esquema propio conservado de la tesis. Cada jaula recibió una botella de agua y una segunda botella con agua, sacarosa o alulosa. La medición compartida se analizó a nivel de jaula. Las secciones de la cohorte se utilizaron para cuantificación POMC/c-FOS y para los análisis gliales cuando existían imágenes y control de calidad adecuados.",
    328: "Tabla 1. Anticuerpos primarios utilizados. Se indican hospedador, dilución, proveedor y catálogo. Las diluciones que variaron entre protocolos se informan en el texto y en los registros de análisis.",
    329: "Tabla 2. Anticuerpos secundarios y DAPI utilizados. La combinación exacta dependió del hospedador del anticuerpo primario y del protocolo montado o flotante.",
    330: "Adquisición de imágenes",
    331: "Las imágenes del hipotálamo mediobasal de la cohorte NPY se adquirieron con Axio Zoom.V16/Apotome.2 y microscopía confocal Zeiss LSM 700 o Leica SP8, según la muestra. Las imágenes se reconstruyeron y analizaron en sus coordenadas nativas. Para visualización se aplicaron transformaciones de intensidad documentadas que no intervinieron en la cuantificación.",
    332: "Las imágenes DAPI/c-FOS/POMC y DAPI/Iba1/GFAP se adquirieron con Keyence BZ-X800. Cada campo bilateral se trató como una sección completa. Las regiones ARC, EM y VMN se delimitaron en cada sección mediante máscaras aceptadas mediante HIL; la segmentación celular se mantuvo automatizada y trazable a sus archivos de origen.",
    333: "Procesamiento de imágenes y delimitación anatómica",
    334: "Los archivos TIFF nativos se asociaron con su animal, sección, condición y canal antes de cuantificar. Las máscaras ARC/EM/VMN se aceptaron mediante HIL en las coordenadas originales y se conservaron junto con recibos de aceptación. DAPI se utilizó para segmentar núcleos; c-FOS, NPY y POMC se asignaron mediante superposición con las regiones celulares definidas en cada flujo. No se utilizaron las imágenes ajustadas para presentación como entrada de la cuantificación.",
    335: "Para la distribución espacial, las posiciones celulares se resumieron en seis capas normalizadas por animal. Este procedimiento permite comparar perfiles relativos, pero no sustituye el registro a un atlas. Los análisis conductuales, histológicos y de procedencia se reconstruyeron mediante reproduce_all_figures.sh; cada figura canónica posee paneles, leyendas y tablas fuente en el repositorio.",
    336: "Análisis de imágenes y estadística",
    337: "Los análisis se ejecutaron con scripts versionados en Python. Cellpose se utilizó para segmentar núcleos y señales celulares en los canales definidos por cada figura. Las proporciones se calcularon primero por imagen y se agregaron por animal. Para Iba1 y GFAP se corrigió el fondo y se normalizó por el área regional. El clasificador de morfología microglial se entrenó con pseudoetiquetas morfométricas; sus clases son propuestas computacionales y no etiquetas expertas de activación biológica.",
    338: "La unidad experimental fue el animal para histología y la jaula para medidas compartidas de bebida o alimento. Se usó ANOVA de una vía como resumen paramétrico y pruebas exactas por permutación o Mann–Whitney por enumeración cuando lo permitió el diseño. Las comparaciones por pares se ajustaron por Holm cuando estaban declaradas como una familia confirmatoria; los perfiles espaciales se ajustaron por Benjamini–Hochberg. No se eliminaron valores señalados por la regla de 1,5 RIC.",
    339: "Los análisis espaciales se resumieron en seis capas por animal y se evaluaron con PERMANOVA exacta y diagnóstico de dispersión. Los tamaños muestrales pequeños, la falta de documentación de aleatorización y cegamiento, el número desigual de secciones y la posible agrupación por jaula se consideraron limitaciones. Cuando la sensibilidad por jaula no respaldó una señal animal, la conclusión se trató como exploratoria. El umbral descriptivo fue p<0,05, pero la interpretación consideró multiplicidad, unidad experimental, tamaño muestral y validez de la medición.",
}


ACKNOWLEDGEMENTS = [
    "Agradezco al Dr. Carlos Farkas Pool por su orientación durante el desarrollo del proyecto y a la Dra. Valentina González Pecchi por su acompañamiento científico y metodológico.",
    "Agradezco al personal del Centro Regional de Estudios Avanzados para la Vida y del Centro de Microscopía Avanzada del Biobío por el apoyo en el cuidado animal, la preparación de muestras y la adquisición de imágenes. También agradezco a quienes participaron en la revisión anatómica mediante HIL y en la organización reproducible de los datos.",
    "Finalmente, agradezco a mi familia, amistades y compañeros por su apoyo durante el programa de Magíster.",
]


SUMMARY_ES = [
    "La alulosa es un epímero de la fructosa con bajo aporte energético. Estudios preclínicos han descrito efectos sobre GLP-1, aferencias vagales y neuronas del núcleo arqueado (ARC), pero la respuesta hipotalámica depende de la vía, la dosis y el contexto de exposición. Este estudio comparó agua, sacarosa al 10 % y alulosa al 13 % en ratones, con énfasis en conducta, c-FOS, neuronas NPY y POMC y marcadores gliales.",
    "En el experimento de una botella, la jaula fue la unidad experimental para consumo y peso. El consumo en el día 6 difirió globalmente entre tratamientos (ANOVA p=0,010; permutación exacta p=0,009), pero ninguna comparación por pares ajustada por Holm alcanzó p<0,05. No se detectó una diferencia global de cambio de peso (ANOVA p=0,119). En el análisis c-FOS/DAPI por animal no hubo diferencias globales en ARC (p=0,133), eminencia media (p=0,466), VMN (p=0,521) ni tejido fuera de las regiones de interés (p=0,163). En ARC, el contraste exacto Sacarosa–Alulosa fue nominalmente significativo antes de ajuste (p=0,0411), pero no se sostuvo en la prueba global ni después de corregir por multiplicidad.",
    "En nueve animales NPY-GFP, c-FOS/DAPI (p=0,095) y la fracción c-FOS+/NPY+ entre núcleos NPY+ (p=0,097) no difirieron globalmente. Esta última fue el resultado de abundancia neuronal más cercano al umbral y su media fue mayor en alulosa que en sacarosa, pero la comparación exacta bilateral fue p=0,100 y no permite concluir que la alulosa induzca hambre. El perfil espacial doble sí difirió entre condiciones (PERMANOVA p=0,0107; q=0,0214), acompañado por heterogeneidad de dispersión (p=0,0143; q=0,0286); por ello, el resultado indica una diferencia multivariada espacial, pero no una separación pura de centroides. En POMC/c-FOS no se detectaron diferencias en ARC o eminencia media.",
    "Tampoco se observaron diferencias globales de señal Iba1 o GFAP en ARC, eminencia media o VMN. Las clases microgliales de la CNN fueron pseudoetiquetas exploratorias. En conjunto, los resultados no apoyan una activación hipotalámica general ni una respuesta POMC o glial específica de alulosa bajo este protocolo. La diferencia con estudios de administración intracerebroventricular o aplicación directa sobre neuronas del ARC puede relacionarse con la exposición oral por pipeta después de adaptación, la señalización gastrointestinal y el aprendizaje previo; esta explicación es una hipótesis que requiere comparación experimental directa.",
    "Palabras clave: alulosa; hipotálamo; núcleo arqueado; NPY; POMC; c-FOS; microglía; astrocitos; conducta alimentaria.",
]


SUMMARY_EN = [
    "Allulose is a low-energy C-3 epimer of fructose. Preclinical studies have reported effects on GLP-1, vagal afferents, and arcuate nucleus (ARC) neurons, but hypothalamic responses depend on administration route, dose, and exposure context. This study compared water, 10% sucrose, and 13% allulose in mice, focusing on behaviour, c-FOS, NPY and POMC neurons, and glial markers.",
    "In the single-bottle experiment, the cage was the experimental unit for consumption and body weight. Day-6 consumption differed globally among treatments (ANOVA p=0.010; exact permutation p=0.009), although no Holm-adjusted pairwise comparison reached p<0.05. Body-weight change did not differ globally (ANOVA p=0.119). Animal-level c-FOS/DAPI showed no global treatment differences in ARC (p=0.133), median eminence (p=0.466), VMN (p=0.521), or tissue outside the regions of interest (p=0.163).",
    "Among nine NPY-GFP animals, neither c-FOS/DAPI (p=0.095) nor the fraction of c-FOS+/NPY+ nuclei among NPY+ nuclei (p=0.097) differed globally. The latter mean was higher with allulose than with sucrose, but the exact two-sided comparison was p=0.100 and does not establish hunger. The double-positive spatial profile differed among conditions (PERMANOVA p=0.0107; q=0.0214) together with heterogeneous dispersion (p=0.0143; q=0.0286); it therefore indicates a multivariate spatial difference but not pure centroid separation. POMC/c-FOS, Iba1 signal, and GFAP signal did not show supported treatment differences.",
    "These results do not support generalized hypothalamic activation or an allulose-specific POMC or glial response under this protocol. Differences from intracerebroventricular administration or direct application to isolated ARC neurons may reflect oral pipette exposure after adaptation, gastrointestinal signalling, and prior learning; this remains a hypothesis requiring a direct route-comparison experiment. The study provides a reproducible analysis framework and defines limits for interpreting small samples and indirect markers.",
    "Keywords: allulose; hypothalamus; arcuate nucleus; NPY; POMC; c-FOS; microglia; astrocytes; feeding behaviour.",
]


ABBREVIATIONS = [
    ("3V", "tercer ventrículo"), ("AgRP", "péptido relacionado con agutí"),
    ("ARC", "núcleo arqueado"), ("BHE", "barrera hematoencefálica"),
    ("CCK", "colecistoquinina"), ("c-FOS", "proteína de respuesta inmediata c-FOS"),
    ("DAPI", "4′,6-diamidino-2-fenilindol"), ("DMT2", "diabetes mellitus tipo 2"),
    ("EM", "eminencia media"), ("ENN", "edulcorante no nutritivo"),
    ("GFAP", "proteína ácida fibrilar glial"), ("GLP-1", "péptido similar al glucagón tipo 1"),
    ("HIL", "human-in-the-loop"), ("Iba1", "molécula adaptadora de unión a calcio ionizado 1"),
    ("IHC", "inmunohistoquímica"), ("ME", "eminencia media, abreviatura usada en paneles"),
    ("NPY", "neuropéptido Y"), ("PBS", "tampón fosfato salino"),
    ("PFA", "paraformaldehído"), ("POMC", "proopiomelanocortina"),
    ("PVH/PVN", "núcleo paraventricular del hipotálamo"),
    ("VMN", "núcleo ventromedial"), ("WT", "tipo silvestre"),
]


RESULT_SECTIONS = [
    (
        "Conducta en el experimento de una botella",
        [
            "El análisis conductual principal se realizó con 12 jaulas (Agua n=3, Sacarosa n=5 y Alulosa n=4). En el día 6, el consumo mostró una diferencia global entre tratamientos (ANOVA F=7,97; p=0,0102; permutación exacta de etiquetas de jaula p=0,00916). Las medias fueron 26,7 mL/jaula para Agua, 107,0 para Sacarosa y 38,8 para Alulosa. Sin embargo, ninguna comparación exacta por pares se mantuvo por debajo de 0,05 después del ajuste de Holm: Agua–Sacarosa p=0,0714, Agua–Alulosa p=0,1429 y Sacarosa–Alulosa p=0,0714.",
            "El cambio medio de peso corporal por jaula en el día 6 no difirió globalmente (ANOVA p=0,119; permutación exacta p=0,0744). Las comparaciones ajustadas por Holm fueron Agua–Sacarosa p=0,514, Agua–Alulosa p=0,514 y Sacarosa–Alulosa p=0,190. La media del grupo Alulosa fue más negativa, pero la variabilidad entre cuatro jaulas fue alta y no permite afirmar un efecto sobre el peso. En consecuencia, el resultado respaldado es una diferencia global de consumo dominada por el alto consumo de sacarosa, no una diferencia confirmada entre alulosa y agua ni un efecto sobre peso.",
        ],
    ),
    (
        "Respuesta c-FOS global en regiones hipotalámicas",
        [
            "La segmentación DAPI/c-FOS se combinó con máscaras anatómicas de ARC, eminencia media (ME), VMN y tejido fuera de las regiones de interés. La proporción c-FOS/DAPI por animal no mostró diferencias globales en ARC (Agua n=4, Sacarosa n=6, Alulosa n=6; ANOVA p=0,133), ME (n=4, 4 y 6; p=0,466), VMN (n=4, 6 y 6; p=0,521) ni tejido residual (n=4, 6 y 6; p=0,163).",
            "En ARC, la comparación exacta Sacarosa–Alulosa sin ajuste fue nominalmente significativa (p=0,0411). El ANOVA global no fue significativo y el ajuste de Holm a través de las cuatro regiones produjo p=0,531 para la prueba global del ARC. Por ello, se conserva y destaca como una señal por pares, pero no se interpreta como evidencia confirmatoria de activación regional. Este análisis reemplaza la versión preliminar de la tesis que informaba p=0,020 y activación específica por alulosa.",
        ],
    ),
    (
        "Respuesta de núcleos NPY positivos y organización espacial",
        [
            "Se analizaron nueve animales NPY-GFP, tres por condición. La proporción de núcleos c-FOS positivos entre núcleos DAPI presentó ANOVA p=0,0951. La fracción de núcleos NPY positivos que también fueron c-FOS positivos presentó ANOVA p=0,0969 y constituyó el resultado de abundancia neuronal más cercano al umbral de significación. Para esta segunda variable, la comparación exacta Sacarosa–Alulosa fue p=0,100, con valores descriptivamente mayores en Alulosa; Agua–Sacarosa y Agua–Alulosa fueron p=0,700. La dirección es compatible con una señal orexigénica relativamente mayor y, por hipótesis, más hambre que con sacarosa, pero no cumple el criterio de significación global y no demuestra que la alulosa active NPY o cause hambre.",
            "El análisis espacial de c-FOS/DAPI en seis capas por animal no difirió entre tratamientos (PERMANOVA exacta p=0,121; q=0,121). En cambio, el perfil espacial de núcleos doble positivos c-FOS+/NPY+ respecto de DAPI mostró una diferencia multivariada (PERMANOVA exacta p=0,0107; q=0,0214). El diagnóstico de dispersión de ese perfil también fue significativo (p=0,0143; q=0,0286). Por tanto, existe una diferencia espacial exploratoria entre condiciones, pero puede incorporar cambios de dispersión, de centroide o de ambos y no define por sí sola una localización anatómica desplazada por alulosa.",
        ],
    ),
    (
        "Respuesta POMC/c-FOS",
        [
            "En la cohorte con inmunomarcación POMC, c-FOS/DAPI no difirió en ME (ANOVA p=0,151) ni ARC (p=0,635). La variable primaria, núcleos doble positivos c-FOS/POMC divididos por el total POMC, tampoco mostró diferencias globales en ME (p=0,512) ni ARC (p=0,430). Las comparaciones exactas y los análisis de sensibilidad por jaula no respaldaron una respuesta específica a alulosa.",
            "La ausencia de señal utilizable en el reportero POMC-GFP del experimento montado se trató como una limitación técnica y no como ausencia biológica de neuronas POMC. La cuantificación final se basó en inmunomarcación POMC en la cohorte disponible y en denominadores por animal.",
        ],
    ),
    (
        "Señal Iba1, GFAP y morfometría microglial",
        [
            "La señal Iba1 corregida por fondo y normalizada por área no difirió entre tratamientos en ARC (ANOVA p=0,919), ME (p=0,828) ni VMN (p=0,585). La señal GFAP normalizada tampoco difirió en ARC (p=0,316), ME (p=0,992) ni VMN (p=0,733). Estos resultados no apoyan un cambio general de estos marcadores gliales asociado a alulosa.",
            "La fracción asignada por la CNN a las pseudoetiquetas morfométricas Activada+Ameboide mostró en ARC un ANOVA por animal p=0,0553 y comparaciones exactas sin ajustar Agua–Sacarosa p=0,0476 y Sacarosa–Alulosa p=0,0260. No obstante, el clasificador no fue entrenado con etiquetas biológicas expertas y la sensibilidad basada en jaulas no confirmó la señal (permutación global por recuentos p=0,125; ANOVA de medias de jaula p=0,103). Se informa como patrón exploratorio del modelo, no como prueba de activación microglial.",
        ],
    ),
    (
        "Análisis complementarios",
        [
            "Los análisis separados por sexo del experimento de una botella tuvieron muy pocas jaulas. En machos, el ANOVA de peso del día 6 fue p=0,049, pero el ómnibus exacto fue p=0,0667 y ninguna comparación ajustada fue significativa. En hembras, el volumen presentó ANOVA p=0,032, pero la permutación exacta global fue p=0,200 y Agua tuvo una sola jaula. No se estimó una interacción tratamiento por sexo y no se infirió dimorfismo sexual.",
            "En la histología H&E de riñón, hígado y bazo, una sola diferencia de composición predicha mantuvo q<0,05: una pseudoetiqueta epitelial hepática para Alulosa frente a Agua (q=0,031). HistoPLUS fue entrenado con tumores humanos y no está validado para tejidos murinos normales; ninguna variable clásica de H&E sobrevivió la corrección global. El resultado se considera una señal computacional generadora de hipótesis.",
            "En el experimento de dos botellas, los ANOVA paramétricos indicaron diferencias globales significativas en consumo de la botella de solución (p=7,17×10⁻⁵), agua (p=0,00379), fracción de la botella de solución (p=1,44×10⁻⁵) y alimento (p=4,58×10⁻⁵), dominadas por la condición Sacarosa. Con solo 2, 3 y 3 jaulas, las comparaciones exactas bilaterales tuvieron p entre 0,100 y 1,000 y no fueron significativas. El cambio de peso no difirió (p=0,661).",
            "En la cohorte de primera exposición después de 16 h de ayuno (Agua n=2, Sacarosa n=3, Alulosa n=3), no se detectaron diferencias en c-FOS/DAPI de ME o ARC ni en c-FOS/ACTH-CLIP de ARC. Los perfiles espaciales fueron no significativos después de ajuste (c-FOS q=0,229; doble positividad q=0,400).",
        ],
    ),
]


DISCUSSION_SECTIONS = [
    (
        "Síntesis de los hallazgos",
        [
            "Este trabajo no confirmó una activación hipotalámica general específica de alulosa. La respuesta c-FOS/DAPI global fue compatible entre tratamientos en ARC, ME, VMN y tejido residual, aunque ARC presentó un contraste nominal Sacarosa–Alulosa sin ajustar (p=0,0411) que no sobrevivió la inferencia global ni la multiplicidad. POMC no presentó diferencias en ninguna de las variables examinadas y los marcadores gliales Iba1 y GFAP tampoco mostraron una respuesta significativa. En conducta, el consumo difirió globalmente, principalmente por el alto consumo de sacarosa, mientras que el peso corporal no mostró un efecto respaldado.",
            "El hallazgo más informativo en la cohorte NPY fue doble. Primero, la fracción c-FOS+/NPY+ fue el resultado de abundancia neuronal más cercano al umbral y mostró una tendencia descriptiva hacia valores mayores con alulosa que con sacarosa, aunque el ANOVA global (p=0,0969) y la comparación exacta (p=0,100) no alcanzaron significación. Esta dirección es compatible con un mayor reclutamiento de neuronas orexigénicas y, por hipótesis, con una señal de hambre relativamente mayor que con sacarosa, pero c-FOS no mide hambre y el estudio no demostró aumento de ingesta inmediatamente después del desafío. Segundo, la organización espacial doble positiva difirió entre condiciones después de corrección por multiplicidad (q=0,0214). La dispersión también difirió (q=0,0286), de modo que el resultado espacial es significativo dentro del análisis declarado, pero no identifica por sí solo un desplazamiento anatómico uniforme.",
        ],
    ),
    (
        "Vía de administración, aprendizaje previo y trabajos de Toshihiko Yada",
        [
            "La comparación con los trabajos de Toshihiko Yada exige separar protocolos que suelen agruparse de manera imprecisa. En el artículo de Nature Communications de Iwasaki et al. (2018), coautorado por Yada, la alulosa se administró por vía peroral hacia el estómago mediante una sonda metálica a ratones en ayuno. Dosis de 1 y 3 g/kg redujeron transitoriamente la ingesta, aumentaron GLP-1 portal y requirieron señalización GLP-1R y aferencias vagales. El efecto no se observó por vía intraperitoneal. Este trabajo apoya una ruta intestinal-vagal, pero no evaluó animales adaptados previamente a beber alulosa ni administró 200 µL en la cavidad oral mediante pipeta.",
            "Yermek et al. (2022a) abordaron una pregunta distinta: aplicaron alulosa directamente en superfusión a neuronas aisladas del ARC y observaron aumentos de calcio en una fracción de neuronas POMC y sensibles a GLP-1. Para el ensayo de conducta usaron inyección intracerebroventricular y detectaron reducción de ingesta. Yermek et al. (2022b) utilizaron el mismo enfoque de exposición directa y comunicaron inhibición de una fracción de neuronas NPY, sensibles a grelina o a baja glucosa; nuevamente, la prueba de alimentación empleó inyección intracerebroventricular. Ninguno de estos protocolos fue intratecal.",
            "Una revisión de Yada, Dezaki e Iwasaki (2025) integró la regulación opuesta por GLP-1 y grelina en islotes, aferencias vagales e hipotálamo. Más recientemente, Rakhat et al. (2026), con Yada como coautor, compararon alulosa con semaglutida oral en ratones con obesidad inducida por dieta: administraron alulosa por gavage a 0,3–5 g/kg y, en experimentos celulares separados, expusieron neuronas aisladas del ARC. Observaron reducción de ingesta y peso, activación de neuronas sensibles a leptina e inhibición de neuronas sensibles a grelina. Este estudio refuerza la importancia de la dosis, el estado metabólico y el uso combinado de observaciones in vivo y exposición celular directa.",
            "Nuestro protocolo difirió en cuatro aspectos con capacidad de modificar la respuesta: exposición repetida a la solución durante la adaptación, dos días de reemplazo por agua, ayuno de 4 h y administración oral de un volumen fijo por pipeta. La pipeta mantiene estimulación orosensorial y no garantiza el mismo depósito gástrico o la misma cinética que una sonda. Además, 200 µL de alulosa al 13 % aportan 26 mg, cerca de 1,0 g/kg para un ratón de 25 g, pero la dosis por kilogramo varía con el peso. Estas diferencias ofrecen una explicación plausible para que la dirección descriptiva de NPY sea opuesta a la inhibición observada con exposición central/directa; no permiten atribuir causalidad sin un estudio factorial de vía y experiencia previa.",
        ],
    ),
    (
        "Interpretación de la tendencia NPY",
        [
            "NPY/AgRP es una población heterogénea cuya actividad suele aumentar durante el ayuno y favorecer la alimentación. Bajo ese marco, una mayor fracción c-FOS+/NPY+ después de alulosa que de sacarosa es compatible con menor supresión de la señal orexigénica. La diferencia puede reflejar que la sacarosa entrega señales postingestivas energéticas rápidas que la alulosa no reproduce, una respuesta condicionada por la experiencia previa o diferencias en señales intestinales y vagales. La hipótesis es coherente con el diseño, pero no está confirmada porque la variable global no fue significativa, n fue tres animales por grupo y no se midió hambre de manera independiente.",
            "El resultado espacial significativo añade evidencia de que las células doble positivas no se distribuyeron de manera idéntica entre condiciones. Sin embargo, los seis anillos se definieron a partir de la nube DAPI de cada campo y no se registraron a un atlas ni al tercer ventrículo. Además, la heterogeneidad de dispersión impide adjudicar la PERMANOVA únicamente a una diferencia de localización promedio. Un estudio de confirmación debería registrar secciones a un atlas, predefinir subregiones del ARC y aumentar el número de animales.",
        ],
    ),
    (
        "Ausencia de una respuesta POMC",
        [
            "POMC no mostró diferencias significativas en ninguno de los aspectos estudiados: c-FOS/DAPI general, fracción c-FOS+/POMC+ en ARC o ME, comparaciones por pares y sensibilidades por jaula. Los perfiles espaciales mostrados fueron descriptivos y no proporcionaron una señal confirmatoria. Por tanto, este estudio no reproduce la activación aguda de POMC descrita después de aplicación directa de alulosa a neuronas aisladas o inyección intracerebroventricular (Yermek et al., 2022a).",
            "La falta de señal POMC no demuestra que la vía melanocortínica sea irrelevante. c-FOS captura una ventana temporal limitada, la inmunodetección identifica solo parte de la población y los tamaños muestrales fueron pequeños. También es posible que una respuesta periférica ocurra en otro intervalo o mediante cambios de actividad que no inducen c-FOS. La conclusión válida se restringe a que no hubo una diferencia detectable con los marcadores, regiones y tiempos utilizados.",
        ],
    ),
    (
        "Respuesta glial",
        [
            "No se detectaron diferencias significativas de señal Iba1 o GFAP en ARC, ME o VMN. En consecuencia, los datos no respaldan activación microglial ni respuesta astrocítica asociada a alulosa bajo este protocolo. GFAP e Iba1 son marcadores parciales y no cubren el repertorio funcional de astrocitos o microglía, pero la ausencia fue consistente en todas las regiones cuantificadas.",
            "Las clases de TinyMorphCNN no cambian esta conclusión. El modelo aprendió pseudoetiquetas derivadas de morfometría y no estados revisados por expertos ni perfiles moleculares. Aunque algunas comparaciones por animal fueron pequeñas, el resultado global del ARC quedó por encima de 0,05 y las sensibilidades por jaula no lo confirmaron. Denominar estas clases 'activadas' es una convención computacional; no constituye evidencia de inflamación.",
        ],
    ),
    (
        "Conducta y unidad experimental",
        [
            "El mayor consumo de sacarosa concuerda con su valor energético y palatabilidad, pero el diseño no separa preferencia individual, derrame y consumo compartido. La jaula fue la unidad experimental correcta porque la botella y el alimento fueron compartidos. Analizar cada ratón como réplica habría inflado artificialmente el tamaño muestral.",
            "Los análisis por sexo y de dos botellas son útiles para describir patrones, pero no permiten conclusiones confirmatorias. En particular, una sola jaula de Agua en hembras impide estimar su variabilidad, y no se probó una interacción tratamiento por sexo. En el experimento de dos botellas, los ANOVA fueron significativos para solución, agua, fracción de solución y alimento, por lo que estas diferencias se informan como resultados paramétricos positivos. Su alcance sigue siendo exploratorio: solo hubo ocho jaulas, los supuestos son difíciles de evaluar y las pruebas exactas bilaterales no alcanzaron significación.",
        ],
    ),
    (
        "Relación con evidencia humana y alcance translacional",
        [
            "La evidencia humana más sólida sobre alulosa se concentra en la atenuación de respuestas posprandiales de glucosa e insulina, no en pérdida de peso ni en activación hipotalámica. El metaanálisis de Osborn et al. (2026) no encontró efectos robustos sobre variables en ayunas, lípidos o composición corporal. Un estudio cruzado reciente con alulosa combinada con estevia no detectó una diferencia de flujo sanguíneo hipotalámico frente a sacarosa, aunque informó efectos exploratorios en regiones de recompensa y vaciamiento gástrico (Smeets et al., 2026). Estos datos no contradicen directamente la tesis, porque difieren en especie, formulación y variable de respuesta.",
            "La guía de la OMS sobre edulcorantes sin azúcar no incluye a la alulosa. No obstante, esa exclusión regulatoria tampoco implica un beneficio clínico. La utilidad de reemplazar sacarosa por alulosa debe juzgarse por reducción neta de azúcares y energía, tolerabilidad gastrointestinal, patrón dietario y evidencia de largo plazo.",
        ],
    ),
    (
        "Fortalezas, limitaciones y trabajo futuro",
        [
            "Entre las fortalezas se encuentran la reconstrucción desde imágenes nativas, la delimitación anatómica HIL documentada, la agregación por animal o jaula, las pruebas exactas cuando fueron posibles, la conservación de valores atípicos y la publicación de paneles, tablas fuente y recibos de procedencia. La revisión final corrigió resultados preliminares y evitó seleccionar comparaciones aisladas.",
            "Las limitaciones principales son los tamaños muestrales pequeños y desiguales, la ausencia de documentación completa sobre aleatorización y cegamiento, la variación de plataformas microscópicas, la falta de una medición directa de hambre, la ausencia de concentraciones circulantes de GLP-1, FGF21 o grelina, y el uso de c-FOS como marcador indirecto. La agrupación por jaula reduce aún más la información independiente en algunas cohortes.",
            "El experimento decisivo debería cruzar animales sin experiencia y adaptados con administración oral por pipeta, gavage e intracerebroventricular a dosis equivalentes, incluir un control de procedimiento y medir ingestión posterior, hormonas periféricas, actividad NPY/POMC en tiempo real y c-FOS en varios intervalos. La vía intracerebroventricular debe reservarse para probar suficiencia central; no reproduce la fisiología de un alimento ingerido.",
        ],
    ),
]


CONCLUSIONS = [
    "La alulosa no produjo una activación c-FOS general confirmada en ARC, ME o VMN bajo el protocolo de exposición oral después de adaptación.",
    "La respuesta POMC no fue significativa en ninguna variable estudiada y la señal glial Iba1/GFAP tampoco mostró diferencias entre tratamientos.",
    "La cohorte NPY presentó el resultado de abundancia neuronal más cercano al umbral: una tendencia descriptiva hacia una mayor fracción c-FOS+/NPY+ con alulosa frente a sacarosa, compatible con una mayor señal orexigénica, pero sin demostrar hambre ni un efecto global significativo.",
    "El perfil espacial c-FOS+/NPY+ sí difirió entre condiciones después de corrección por multiplicidad; la heterogeneidad de dispersión exige interpretarlo como una diferencia multivariada espacial sin asignar un desplazamiento anatómico único.",
    "Las diferencias respecto de los trabajos de Yada y colaboradores pueden relacionarse con administración oral, aprendizaje previo, estado metabólico y dosis, pero esta explicación debe probarse mediante comparación directa de vías.",
    "Los resultados conductuales indican mayor consumo de sacarosa; no confirman una diferencia alulosa–agua ni un efecto sobre el peso corporal.",
    "En el experimento suplementario de dos botellas, cuatro desenlaces conductuales presentaron ANOVA global significativo, pero la muestra de ocho jaulas y las pruebas exactas no significativas limitan la inferencia confirmatoria.",
]


MAIN_FIGURES = [
    (
        "Figura 1. Conducta en el experimento de una botella",
        [ROOT / "Fig1" / "panels" / f"Figure_1_behavior_experiment_1_Panel_{letter}_spanish.png"
         for letter in "ABCD"],
        ROOT / "Fig1" / "legends" / "Figure_1_behavior_experiment_1_LEGEND_spanish.txt",
        "Figura 1",
    ),
    (
        "Figura 2. Cuantificación global de c-FOS en regiones hipotalámicas",
        [
            ROOT / "Fig2" / "panels" / "Panel_A_experimental_timeline_spanish.png",
            ROOT / "Fig2" / "panels" / "Panel_B_human_regions_spanish.png",
            ROOT / "Fig2" / "panels" / "Panel_C_native_microscopy_spanish.png",
            ROOT / "Fig2" / "panels" / "Panel_D_animal_statistics_spanish.png",
        ],
        ROOT / "Fig2" / "legends" / "Figure_2_legend_spanish.txt",
        "Figura 2",
    ),
    (
        "Figura 3. Respuesta c-FOS/NPY y organización espacial",
        [ROOT / "Fig3" / "panels" / f"Figure_3_cFos_NPY_Panel_{letter}_spanish.png"
         for letter in "ABCDEFG"],
        ROOT / "Fig3" / "legends" / "Figure_3_cFos_NPY_LEGEND_spanish.txt",
        "Figura 3",
    ),
    (
        "Figura 4. Respuesta POMC/c-FOS",
        [
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_A_dapi_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_B_cfos_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_C_pomc_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_D_cellpose_cartoon_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_E_microscopy_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_F_microscopy_npy_gfp_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_G_cfos_dapi_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_H_pomc_activation_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_I_spatial_cfos_spanish.png",
            ROOT / "Fig4" / "panels" / "Figure_ARC_ME_panel_J_spatial_cfos_pomc_spanish.png",
        ],
        ROOT / "Fig4" / "legends" / "Figure_4_LEGEND_spanish.txt",
        "Figura 4",
    ),
    (
        "Figura 5. Respuesta glial y morfometría microglial",
        [ROOT / "Fig5" / "panels" /
         f"Figure_GFAP_Iba1_microglia_ARC_ME_Panel_{letter}_spanish.png"
         for letter in "ABCDEFGHI"],
        ROOT / "Fig5" / "legends" / "Figure_GFAP_Iba1_microglia_ARC_ME_LEGEND_spanish.txt",
        "Figura 5",
    ),
]


SUPPLEMENTARY_FIGURES = [
    (
        "Figura suplementaria S1. Experimento de una botella en machos",
        [ROOT / "FigS1" / "panels" /
         f"Figure_S1_single_bottle_male_Panel_{letter}_spanish.png"
         for letter in "ABCD"],
        ROOT / "FigS1" / "legends" / "Figure_S1_single_bottle_male_LEGEND_spanish.txt",
        "Figura S1",
    ),
    (
        "Figura suplementaria S2. Experimento de una botella en hembras",
        [ROOT / "FigS2" / "panels" /
         f"Figure_S2_single_bottle_female_Panel_{letter}_spanish.png"
         for letter in "ABCD"],
        ROOT / "FigS2" / "legends" / "Figure_S2_single_bottle_female_LEGEND_spanish.txt",
        "Figura S2",
    ),
    (
        "Figura suplementaria S3. Histología H&E multiorgánica",
        [ROOT / "FigS3" / "panels" / f"Figure_S3_Panel_{letter}_spanish.png"
         for letter in "ABCDEF"],
        ROOT / "FigS3" / "legends" / "Figure_S3_LEGEND_spanish.txt",
        "Figura S3",
    ),
    (
        "Figura suplementaria S4. Experimento de dos botellas",
        [
            ROOT / "FigS4" / "panels" / "Figure_S4_behavior_preference_Panel_A_solution_spanish.png",
            ROOT / "FigS4" / "panels" / "Figure_S4_behavior_preference_Panel_B_water_spanish.png",
            ROOT / "FigS4" / "panels" / "Figure_S4_behavior_preference_Panel_C_solution_fraction_spanish.png",
            ROOT / "FigS4" / "panels" / "Figure_S4_behavior_preference_Panel_D_food_spanish.png",
            ROOT / "FigS4" / "panels" / "Figure_S4_behavior_preference_Panel_E_body_weight_spanish.png",
        ],
        ROOT / "FigS4" / "legends" / "Figure_S4_behavior_preference_LEGEND_spanish.txt",
        "Figura S4",
    ),
    (
        "Figura suplementaria S5. Primera exposición después de ayuno",
        [ROOT / "FigS5" / "panels" / f"Figure_S5_panel_{letter}_spanish.png"
         for letter in "ABCDEFGHIJ"],
        ROOT / "FigS5" / "legends" / "Figure_S5_LEGEND_spanish.txt",
        "Figura S5",
    ),
]


THESIS_FIGURE_START = 9

PAPER_FIGURES = [
    (
        "Figure 1",
        ROOT / "Fig1" / "Figure_1_behavior_experiment_1.png",
        ROOT / "Fig1" / "legends" / "Figure_1_behavior_experiment_1_LEGEND.txt",
    ),
    (
        "Figure 2",
        ROOT / "Fig2" / "Figure_2.png",
        ROOT / "Fig2" / "legends" / "Figure_2_legend.txt",
    ),
    (
        "Figure 3",
        ROOT / "Fig3" / "Figure_3_cFos_NPY.png",
        ROOT / "Fig3" / "legends" / "Figure_3_cFos_NPY_LEGEND.txt",
    ),
    (
        "Figure 4",
        ROOT / "Fig4" / "Figure_4.png",
        ROOT / "Fig4" / "legends" / "Figure_4_LEGEND.txt",
    ),
    (
        "Figure 5",
        ROOT / "Fig5" / "Figure_5.png",
        ROOT / "Fig5" / "legends" / "Figure_GFAP_Iba1_microglia_ARC_ME_LEGEND.txt",
    ),
    (
        "Figure S1",
        ROOT / "FigS1" / "Figure_S1_single_bottle_male.png",
        ROOT / "FigS1" / "legends" / "Figure_S1_single_bottle_male_LEGEND.txt",
    ),
    (
        "Figure S2",
        ROOT / "FigS2" / "Figure_S2_single_bottle_female.png",
        ROOT / "FigS2" / "legends" / "Figure_S2_single_bottle_female_LEGEND.txt",
    ),
    (
        "Figure S3",
        ROOT / "FigS3" / "Figure_S3.png",
        ROOT / "FigS3" / "legends" / "Figure_S3_LEGEND.txt",
    ),
    (
        "Figure S4",
        ROOT / "FigS4" / "Figure_S4_behavior_preference.png",
        ROOT / "FigS4" / "legends" / "Figure_S4_behavior_preference_LEGEND.txt",
    ),
    (
        "Figure S5",
        ROOT / "FigS5" / "Figure_S5.png",
        ROOT / "FigS5" / "legends" / "Figure_S5_LEGEND.txt",
    ),
]


S5_PANEL_CAPTIONS = {
    "A": "Canal DAPI del campo representativo registrado FR4-1 de la condición Alulosa después de 16 horas de ayuno y primera exposición a la solución.",
    "B": "Canal c-FOS de los mismos píxeles registrados del campo representativo FR4-1.",
    "C": "Canal ACTH/CLIP de los mismos píxeles registrados del campo representativo FR4-1.",
    "D": "Clases de segmentación Cellpose centradas en núcleos, con anatomía HIL de ARC, ME y VMN basada en DAPI.",
    "E": "Composición registrada exacta de DAPI, c-FOS y ACTH/CLIP; el recuadro señala una célula representativa.",
    "F": "Microscopía que conserva la intensidad natural de ACTH/CLIP dentro de las regiones celulares aceptadas.",
    "G": "Núcleos c-FOS positivos divididos por núcleos DAPI en ME y ARC; las barras muestran media y desviación estándar por adquisición.",
    "H": "En ARC, núcleos dobles c-FOS/ACTH-CLIP divididos por el total de núcleos ACTH/CLIP positivos.",
    "I": "Mapa de ocurrencia espacial y perfil de seis capas por animal para núcleos c-FOS positivos.",
    "J": "Mapa de ocurrencia espacial y perfil de seis capas por animal para núcleos dobles c-FOS/ACTH-CLIP positivos.",
}


def panel_legend_file(group_legend: Path, letter: str) -> Path | None:
    """Return the canonical Spanish legend for one exported panel, if present."""
    needle = f"panel_{letter.lower()}"
    candidates = []
    for path in group_legend.parent.glob("*.txt"):
        normalized = path.name.casefold()
        if needle in normalized and "spanish" in normalized and "legend" in normalized:
            candidates.append(path)
    return sorted(candidates, key=lambda path: path.name.casefold())[0] if candidates else None


def clean_panel_description(text: str, letter: str) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    patterns = (
        rf"^Figura\s+[^.]*,\s*panel\s+{letter}\.\s*",
        rf"^Figura\s+[^.]*\.\s*\({letter}\)\s*",
        rf"^Panel\s+{letter}\.\s*",
        rf"^\({letter}\)\s*",
    )
    for pattern in patterns:
        updated = re.sub(pattern, "", cleaned, count=1, flags=re.IGNORECASE)
        if updated != cleaned:
            cleaned = updated
            break
    return cleaned.strip()


def panel_description(group_legend: Path, letter: str) -> str:
    """Extract a standalone caption from the canonical Spanish legend."""
    if group_legend.name == "Figure_S5_LEGEND_spanish.txt":
        return S5_PANEL_CAPTIONS[letter]

    individual = panel_legend_file(group_legend, letter)
    if individual:
        blocks = [block.strip() for block in re.split(
            r"\n\s*\n", individual.read_text(encoding="utf-8")) if block.strip()]
        marker = re.compile(
            rf"(?:\({letter}\)|panel[_ ]{letter}\b)", re.IGNORECASE)
        for block in blocks:
            if marker.search(block):
                return clean_panel_description(block, letter)

    full = group_legend.read_text(encoding="utf-8")
    match = re.search(
        rf"^\({letter}\)\s*(.*?)(?=^\([A-Z]\)\s|\Z)",
        full,
        flags=re.MULTILINE | re.DOTALL,
    )
    if match:
        return clean_panel_description(match.group(1), letter)
    raise RuntimeError(
        f"No se pudo extraer la leyenda del panel {letter}: {group_legend}")


def group_general_note(group_legend: Path) -> str:
    """Keep statistical/methodological notes that apply to a complete panel set."""
    blocks = [re.sub(r"\s+", " ", block.strip())
              for block in re.split(
                  r"\n\s*\n", group_legend.read_text(encoding="utf-8"))
              if block.strip()]
    notes = []
    for position, block in enumerate(blocks):
        if position == 0:
            continue
        if re.match(r"^\([A-Z](?:[-–][A-Z])?\)", block):
            continue
        if re.match(r"^Panel\s+[A-Z]\.", block, flags=re.IGNORECASE):
            continue
        if (group_legend.name == "Figure_S5_LEGEND_spanish.txt" and
                "A-C, canales" in block):
            continue
        notes.append(block)
    return " ".join(notes)


def recode_panel_references(text: str, first_number: int) -> str:
    """Replace paper-panel references with thesis figure numbers."""
    def number(letter: str) -> int:
        return first_number + ord(letter.upper()) - ord("A")

    substitutions = (
        (r"\bdel panel\s+([A-J])\b",
         lambda m: f"de la Figura {number(m.group(1))}"),
        (r"\ben el panel\s+([A-J])\b",
         lambda m: f"en la Figura {number(m.group(1))}"),
        (r"\blos paneles\s+([A-J]),\s*([A-J])\s+y\s+([A-J])\b",
         lambda m: (f"las Figuras {number(m.group(1))}, "
                    f"{number(m.group(2))} y {number(m.group(3))}")),
        (r"\bpaneles\s+([A-J])\s*[/–-]\s*([A-J])\b",
         lambda m: f"Figuras {number(m.group(1))}–{number(m.group(2))}"),
        (r"\bpanel(?:es)?\s+([A-J])\s+y\s+([A-J])\b",
         lambda m: f"Figuras {number(m.group(1))} y {number(m.group(2))}"),
        (r"\bpanel(?:es)?\s+([A-J])\b",
         lambda m: f"Figura {number(m.group(1))}"),
        (r"\bperfiles\s+([A-J])\s*[–-]\s*([A-J])\b",
         lambda m: (f"perfiles de las Figuras {number(m.group(1))}–"
                    f"{number(m.group(2))}")),
    )
    converted = text
    for pattern, replacement in substitutions:
        converted = re.sub(
            pattern, replacement, converted, flags=re.IGNORECASE)
    if first_number == 24:
        converted = converted.replace(
            "de A, B y C", "de las Figuras 24, 25 y 26")
        converted = converted.replace(
            "produce exactamente E", "produce exactamente la Figura 28")
    return converted


def thesis_panel_records() -> list[dict]:
    records = []
    number = THESIS_FIGURE_START
    groups = MAIN_FIGURES + SUPPLEMENTARY_FIGURES
    for group_index, (title, panels, legend, paper_prefix) in enumerate(groups):
        group_first = number
        for panel_index, panel in enumerate(panels):
            letter = chr(ord("A") + panel_index)
            records.append({
                "number": number,
                "key": f"fig_{number}",
                "group_index": group_index,
                "group_title": title,
                "paper_prefix": paper_prefix,
                "letter": letter,
                "panel": panel,
                "legend": legend,
                "caption": recode_panel_references(
                    panel_description(legend, letter), group_first),
            })
            number += 1
    return records


YADA_2025_REFERENCE = {
    "type": "article-journal",
    "title": "GLP-1 and ghrelin inversely regulate insulin secretion and action in pancreatic islets, vagal afferents, and hypothalamus for controlling glycemia and feeding",
    "author": [
        {"family": "Yada", "given": "Toshihiko"},
        {"family": "Dezaki", "given": "Katsuya"},
        {"family": "Iwasaki", "given": "Yusaku"},
    ],
    "issued": {"date-parts": [[2025, 6, 1]]},
    "container-title": "American Journal of Physiology-Cell Physiology",
    "volume": "328",
    "issue": "6",
    "page": "C1793–C1807",
    "DOI": "10.1152/ajpcell.00168.2025",
    "URL": "https://pubmed.ncbi.nlm.nih.gov/40241252/",
}


def prepare_preserved_media() -> None:
    """Materialize preserved thesis-only gels from the source DOCX in /tmp."""
    MEDIA.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(SOURCE) as archive:
        for filename in PRESERVED_MEDIA_NAMES:
            member = f"word/media/{filename}"
            try:
                payload = archive.read(member)
            except KeyError as exc:
                raise FileNotFoundError(
                    f"Falta el recurso preservado {member} en {SOURCE}"
                ) from exc
            target = MEDIA / filename
            if target.is_file() and target.read_bytes() == payload:
                continue
            temporary = target.with_name(target.name + ".building")
            temporary.write_bytes(payload)
            temporary.replace(target)

def validate_inputs() -> list[Path]:
    prepare_preserved_media()
    required = [SOURCE, *[legend for _, _, legend, _ in MAIN_FIGURES],
                *[legend for _, _, legend, _ in SUPPLEMENTARY_FIGURES]]
    panels = [panel for _, group, _, _ in MAIN_FIGURES + SUPPLEMENTARY_FIGURES
              for panel in group]
    required.extend(panels)
    required.extend([MEDIA / "image27.png", MEDIA / "image28.png"])
    for _, figure, legend in PAPER_FIGURES:
        required.extend([figure, legend])
    missing = [path for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Faltan entradas: " + ", ".join(map(str, missing)))
    if len(panels) != 63 or len(set(panels)) != 63:
        raise RuntimeError(f"Se esperaban 63 paneles españoles únicos; se hallaron {len(set(panels))}")
    for panel in panels:
        with Image.open(panel) as image:
            image.verify()
    records = thesis_panel_records()
    if [record["number"] for record in records] != list(range(9, 72)):
        raise RuntimeError("La numeración de figuras de tesis debe abarcar 9–71 sin saltos")
    return panels


def load_page_map(path: Path | None = None) -> dict[str, int | str]:
    source = path or PAGE_MAP_PATH
    if not source.is_file():
        return {}
    data = json.loads(source.read_text(encoding="utf-8"))
    return {str(key): value for key, value in data.items()}


def index_line(label: str, key: str, page_map: dict,
               width: int = 88) -> str:
    page = str(page_map.get(key, "—"))
    dots = "." * max(4, width - len(label) - len(page))
    return f"{label} {dots} {page}"


def short_caption(text: str, limit: int = 86) -> str:
    first = re.split(r"(?<=[.!?])\s+", re.sub(r"\s+", " ", text).strip())[0]
    first = first.rstrip(".")
    if len(first) <= limit:
        return first
    return first[:limit - 1].rstrip() + "…"


def set_compact_index(paragraph: Paragraph, lines: list[str],
                      size: float = BODY_FONT_SIZE) -> None:
    set_text(paragraph, "\n".join(lines), align=WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, size=Pt(size))


def fill_front_matter(document: Document, paragraphs: list[Paragraph],
                      page_map: dict[str, int | str]) -> None:
    set_text(paragraphs[131], "ÍNDICE", bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, size=TITLE_FONT_SIZE)
    set_text(paragraphs[157], "ÍNDICE DE FIGURAS", bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, size=TITLE_FONT_SIZE)
    set_text(paragraphs[183], "ÍNDICE DE TABLAS", bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, size=TITLE_FONT_SIZE)
    for index in (55, 81, 105, 209):
        set_text(paragraphs[index], paragraphs[index].text, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, size=TITLE_FONT_SIZE)

    for offset, text_value in enumerate(ACKNOWLEDGEMENTS, 56):
        set_text(paragraphs[offset], text_value, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for offset, text_value in enumerate(SUMMARY_ES, 82):
        set_text(paragraphs[offset], text_value, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for offset, text_value in enumerate(SUMMARY_EN, 106):
        set_text(paragraphs[offset], text_value, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    contents_spec = [
        ("AGRADECIMIENTOS", "acknowledgements"),
        ("RESUMEN", "summary"),
        ("ABSTRACT", "abstract"),
        ("ÍNDICE DE FIGURAS", "figure_index"),
        ("ÍNDICE DE TABLAS", "table_index"),
        ("ABREVIATURAS", "abbreviations"),
        ("I. INTRODUCCIÓN", "introduction"),
        ("II. HIPÓTESIS Y OBJETIVOS", "hypothesis"),
        ("III. METODOLOGÍA", "methodology"),
        ("IV. RESULTADOS", "results"),
        ("V. DISCUSIÓN", "discussion"),
        ("VI. CONCLUSIONES", "conclusions"),
        ("VII. REFERENCIAS", "references"),
        ("ANEXO. MANUSCRITO EN PREPARACIÓN PARA ENVÍO", "paper_appendix"),
    ]
    set_compact_index(
        paragraphs[132],
        [index_line(label, key, page_map) for label, key in contents_spec],
        size=BODY_FONT_SIZE,
    )
    figure_spec = [
        ("Figura 1. Núcleos hipotalámicos", "fig_1"),
        ("Figura 2. Mapa celular HypoMap", "fig_2"),
        ("Figura 3. Sistema melanocortínico", "fig_3"),
        ("Figura 4. Estructura de fructosa y alulosa", "fig_4"),
        ("Figura 5. Genotipificación POMC-GFP", "fig_5"),
        ("Figura 6. Genotipificación NPY-GFP", "fig_6"),
        ("Figura 7. Diseño de una botella", "fig_7"),
        ("Figura 8. Diseño de dos botellas", "fig_8"),
    ]
    figure_spec.extend(
        (f"Figura {record['number']}. {short_caption(record['caption'])}",
         record["key"])
        for record in thesis_panel_records()
    )
    for paper_label, _, _ in PAPER_FIGURES:
        key_suffix = paper_label.replace("Figure ", "").casefold()
        figure_spec.append(
            (f"Anexo del artículo — {paper_label}",
             f"paper_fig_{key_suffix}"))
    set_compact_index(
        paragraphs[158],
        [index_line(label, key, page_map) for label, key in figure_spec],
        size=BODY_FONT_SIZE,
    )
    table_lines = [
        index_line("Tabla 1. Anticuerpos primarios", "table_1", page_map),
        index_line("Tabla 2. Anticuerpos secundarios y DAPI", "table_2", page_map),
    ]
    set_compact_index(paragraphs[184], table_lines, size=BODY_FONT_SIZE)
    abbreviations = "\n".join(f"{short}: {long}" for short, long in ABBREVIATIONS)
    set_text(paragraphs[210], abbreviations, align=WD_ALIGN_PARAGRAPH.LEFT)


def apply_source_revisions(document: Document, paragraphs: list[Paragraph]) -> None:
    for index, text_value in {**INTRO_REPLACEMENTS, **METHOD_REPLACEMENTS}.items():
        set_text(paragraphs[index], text_value, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for index in (211, 269, 293):
        apply_heading_style(paragraphs[index], 1)
    heading_2_indices = {
        214, 216, 221, 235, 242, 246, 249, 253, 259, 262, 264,
        270, 272, 274, 294, 296, 298, 300, 302, 304, 306, 308, 310,
        314, 316, 323, 330, 333, 336,
    }
    for index in heading_2_indices:
        apply_heading_style(paragraphs[index], 2)

    caption_indices = {
        219, 231, 232, 240, 241, 256, 257,
        321, 322, 326, 327, 328, 329,
    }
    for index in caption_indices:
        paragraphs[index].style = document.styles["Caption"]
        paragraphs[index].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraphs[index].paragraph_format.line_spacing = THESIS_LINE_SPACING
        for run in paragraphs[index].runs:
            set_run_font(run, size=Pt(BODY_FONT_SIZE))

    clear_paragraph(paragraphs[320])
    timeline = ROOT / "Fig2" / "panels" / "Panel_A_experimental_timeline_spanish.png"
    width, height = picture_dimensions(timeline, max_width=6.0, max_height=4.2)
    paragraphs[320].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[320].add_run().add_picture(str(timeline), width=Inches(width), height=Inches(height))

    anchor = paragraphs[315]
    for image_name, caption_text in (
        ("image27.png", "Figura 5. Gel de genotipificación POMC-GFP conservado de la tesis original."),
        ("image28.png", "Figura 6. Gel de genotipificación NPY-GFP conservado de la tesis original."),
    ):
        image_paragraph = insert_after(anchor)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        path = MEDIA / image_name
        width, height = picture_dimensions(path, max_width=5.8, max_height=3.4)
        image_paragraph.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
        caption_paragraph = insert_after(image_paragraph, caption_text)
        caption_paragraph.style = document.styles["Caption"]
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        caption_paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
        for run in caption_paragraph.runs:
            set_run_font(run, size=Pt(BODY_FONT_SIZE))
        anchor = caption_paragraph

    primary_table, secondary_table = document.tables[:2]
    set_cell_text(primary_table.rows[0].cells[0], "Anticuerpo primario")
    set_cell_text(primary_table.rows[0].cells[1], "Hospedador")
    secondary_rows = [
        ("Anticuerpo secundario", "Especie reconocida", "Dilución", "Empresa", "Catálogo"),
        ("Alexa Fluor 488", "Mouse", "1:200", "Cell Signaling", "4408"),
        ("Alexa Fluor 647", "Rabbit", "1:200", "Jackson ImmunoResearch", "711-605-152"),
        ("Alexa Fluor 647", "Guinea pig", "1:500", "Jackson ImmunoResearch", "706-605-148"),
        ("Alexa Fluor 555", "Goat", "1:500", "Thermo Fisher Scientific", "A-21432"),
        ("Alexa Fluor 647", "Rabbit", "1:500", "Thermo Fisher Scientific", "A-31573"),
        ("DAPI", "-", "1:1.000", "-", "-"),
    ]
    for row, values in zip(secondary_table.rows, secondary_rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)


def author_sort_name(record: dict) -> str:
    authors = record.get("author") or []
    if not authors:
        return record.get("title", "")
    return authors[0].get("family") or authors[0].get("literal") or ""


def author_text(record: dict) -> str:
    names = []
    for author in record.get("author") or []:
        if author.get("literal"):
            names.append(author["literal"])
            continue
        family = author.get("family", "")
        initials = " ".join(
            f"{part[0]}." for part in
            re.findall(r"[A-Za-zÀ-ÿ]+", author.get("given", "")))
        names.append(f"{family}, {initials}".strip().rstrip(","))
    if not names:
        return "Autor no indicado"
    if len(names) > 20:
        names = names[:19] + ["…", names[-1]]
    if len(names) == 1:
        return names[0]
    if "…" in names:
        return ", ".join(names)
    return ", ".join(names[:-1]) + ", & " + names[-1]


def issued_year(record: dict) -> str:
    parts = ((record.get("issued") or {}).get("date-parts") or [["s. f."]])[0]
    return str(parts[0]) if parts else "s. f."


def reference_key(record: dict) -> str:
    doi = str(record.get("DOI", "")).lower().strip()
    if doi:
        return "doi:" + doi
    return "title:" + re.sub(r"\W+", "", record.get("title", "").lower())


def load_references() -> list[dict]:
    records = []
    for path in (LITERATURE / "zotero_items.json",
                 LITERATURE / "current_literature_additions.json"):
        records.extend(json.loads(path.read_text(encoding="utf-8")))
    records.append(YADA_2025_REFERENCE)
    unique = {}
    for record in records:
        unique[reference_key(record)] = record
    return sorted(unique.values(), key=lambda item: (
        author_sort_name(item).casefold(), issued_year(item), item.get("title", "").casefold()))


def normalize_citation_name(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("’", "'")
    value = unicodedata.normalize("NFKD", value)
    value = "".join(character for character in value
                    if not unicodedata.combining(character))
    value = re.sub(r"\s+", " ", value).strip().casefold()
    return value


class CitationManager:
    """Map author-year prose to bracketed numbers in first-citation order."""

    def __init__(self, records: list[dict]):
        self.records = records
        self.by_key = {reference_key(record): record for record in records}
        self.by_author_year: dict[tuple[str, str], dict] = {}
        for record in records:
            key = (normalize_citation_name(author_sort_name(record)),
                   issued_year(record))
            if key in self.by_author_year:
                raise RuntimeError(
                    f"Referencia autor-año ambigua para {key[0]} {key[1]}")
            self.by_author_year[key] = record
        self.cited_keys: list[str] = []
        self.number_by_key: dict[str, int] = {}

    @staticmethod
    def author_from_label(label: str) -> str:
        label = re.sub(r"^Adaptado de\s+", "", label.strip(),
                       flags=re.IGNORECASE)
        label = re.sub(r"\s+et al\.$", "", label).strip()
        if " y " in label:
            label = label.split(" y ", 1)[0]
        return label

    def record_for(self, author: str, year: str) -> dict | None:
        author = self.author_from_label(author)
        year = re.sub(r"[a-z]$", "", year.strip(), flags=re.IGNORECASE)
        return self.by_author_year.get(
            (normalize_citation_name(author), year))

    def citation_for(self, author: str, year: str) -> str | None:
        record = self.record_for(author, year)
        if record is None:
            return None
        key = reference_key(record)
        if key not in self.number_by_key:
            self.cited_keys.append(key)
            self.number_by_key[key] = len(self.cited_keys)
        return f"[{self.number_by_key[key]}]"

    def convert(self, text: str) -> str:
        if not text or not re.search(r"(?:19|20)\d{2}", text):
            return text

        # Correct the two distinct 2022 Yada-group papers before rendering:
        # Yermek is the POMC paper and Rakhat is the NPY paper.
        combined = "Yermek et al. (2022a, 2022b)"
        if combined in text:
            first = self.citation_for("Yermek", "2022")
            second = self.citation_for("Rakhat", "2022")
            if first and second:
                text = text.replace(
                    combined,
                    f"Yermek et al. {first} y Rakhat et al. {second}")

        special_narrative = (
            (r"Yermek et al\. \(2022a\)", "Yermek et al.", "Yermek", "2022"),
            (r"Yermek et al\. \(2022b\)", "Rakhat et al.", "Rakhat", "2022"),
            (r"Yada, Dezaki e Iwasaki \(2025\)",
             "Yada, Dezaki e Iwasaki", "Yada", "2025"),
        )
        for pattern, visible, author, year in special_narrative:
            if re.search(pattern, text) is None:
                continue
            citation = self.citation_for(author, year)
            if citation:
                text = re.sub(pattern, f"{visible} {citation}", text)

        narrative_pattern = re.compile(
            r"(?P<author>[A-ZÁÉÍÓÚÜÑ][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+) "
            r"et al\. \((?P<year>(?:19|20)\d{2})\)")

        def replace_narrative(match: re.Match) -> str:
            citation = self.citation_for(
                match.group("author"), match.group("year"))
            if citation is None:
                return match.group(0)
            return f"{match.group('author')} et al. {citation}"

        text = narrative_pattern.sub(replace_narrative, text)

        parenthetical_pattern = re.compile(
            r"\((?P<inside>[^()\n]{0,360}(?:19|20)\d{2}[a-z]?"
            r"[^()\n]{0,360})\)")

        def replace_parenthetical(match: re.Match) -> str:
            inside = match.group("inside").strip()
            citations: list[str] = []
            for segment in [piece.strip() for piece in inside.split(";")]:
                parsed = re.fullmatch(
                    r"(?P<author>.+?),\s*(?P<years>(?:19|20)\d{2}[a-z]?"
                    r"(?:\s*,\s*(?:19|20)\d{2}[a-z]?)*)",
                    segment)
                if parsed is None:
                    return match.group(0)
                segment_citations = []
                for year in re.findall(
                        r"(?:19|20)\d{2}[a-z]?", parsed.group("years")):
                    citation = self.citation_for(
                        parsed.group("author"), year)
                    if citation is None:
                        return match.group(0)
                    segment_citations.append(citation)
                citations.extend(segment_citations)
            citations = list(dict.fromkeys(citations))
            return " ".join(citations) if citations else match.group(0)

        return parenthetical_pattern.sub(replace_parenthetical, text)

    def reference_records(self) -> list[dict]:
        cited = [self.by_key[key] for key in self.cited_keys]
        cited_set = set(self.cited_keys)
        uncited = [record for record in self.records
                   if reference_key(record) not in cited_set]
        return cited + uncited


def format_reference(record: dict) -> str:
    authors = author_text(record)
    year = issued_year(record)
    title = re.sub(r"\s+", " ", record.get("title", "Sin título")).strip().rstrip(".")
    container = record.get("container-title") or record.get("publisher") or ""
    volume = str(record.get("volume", "")).strip()
    issue = str(record.get("issue", "")).strip()
    pages = str(record.get("page", "")).strip()
    venue = container
    if volume:
        venue += f", {volume}"
    if issue:
        venue += f"({issue})"
    if pages:
        venue += f", {pages}"
    doi = str(record.get("DOI", "")).strip()
    url = str(record.get("URL", "")).strip()
    locator = f"https://doi.org/{doi}" if doi else url
    parts = [f"{authors} ({year}). {title}."]
    if venue:
        parts.append(venue.rstrip(".") + ".")
    if locator:
        parts.append(locator)
    return " ".join(parts)


def add_references(document: Document, records: list[dict]) -> int:
    document.add_page_break()
    add_heading(document, "VII. REFERENCIAS", 1)
    for number, record in enumerate(records, 1):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        paragraph.paragraph_format.line_spacing = THESIS_LINE_SPACING
        paragraph.paragraph_format.space_after = Pt(6)
        number_run = paragraph.add_run(f"[{number}] ")
        number_run.bold = True
        set_run_font(number_run, size=Pt(BODY_FONT_SIZE))
        authors = author_text(record)
        year = issued_year(record)
        title = re.sub(
            r"\s+", " ", record.get("title", "Sin título")).strip().rstrip(".")
        lead = paragraph.add_run(f"{authors} ({year}). {title}. ")
        set_run_font(lead, size=Pt(BODY_FONT_SIZE))

        container = record.get("container-title") or record.get("publisher") or ""
        volume = str(record.get("volume", "")).strip()
        issue = str(record.get("issue", "")).strip()
        pages = str(record.get("page", "")).strip()
        if container:
            venue = paragraph.add_run(str(container).rstrip("."))
            venue.italic = True
            set_run_font(venue, size=Pt(BODY_FONT_SIZE))
            if volume:
                comma = paragraph.add_run(", ")
                set_run_font(comma, size=Pt(BODY_FONT_SIZE))
                volume_run = paragraph.add_run(volume)
                volume_run.italic = True
                set_run_font(volume_run, size=Pt(BODY_FONT_SIZE))
            if issue:
                issue_run = paragraph.add_run(f"({issue})")
                set_run_font(issue_run, size=Pt(BODY_FONT_SIZE))
            if pages:
                pages_run = paragraph.add_run(f", {pages}")
                set_run_font(pages_run, size=Pt(BODY_FONT_SIZE))
            stop = paragraph.add_run(". ")
            set_run_font(stop, size=Pt(BODY_FONT_SIZE))

        doi = str(record.get("DOI", "")).strip()
        url = str(record.get("URL", "")).strip()
        locator = f"https://doi.org/{doi}" if doi else url
        if locator:
            locator_run = paragraph.add_run(locator)
            set_run_font(locator_run, size=Pt(BODY_FONT_SIZE))
    return len(records)


def add_paper_appendix(document: Document) -> None:
    """Append the ten final English master figures exactly as paper artefacts."""
    document.add_page_break()
    add_heading(document, "ANEXO. MANUSCRITO EN PREPARACIÓN PARA ENVÍO", 1)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run(PAPER_TITLE)
    title_run.bold = True
    set_run_font(title_run, size=Pt(TITLE_FONT_SIZE))

    authors = document.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(8)
    author_run = authors.add_run(", ".join(PAPER_AUTHORS))
    set_run_font(author_run, size=Pt(BODY_FONT_SIZE))

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.line_spacing = THESIS_LINE_SPACING
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run(
        "* Autor de correspondencia. El orden indicado corresponde a la lista entregada "
        "para esta versión. Las afiliaciones y la forma bibliográfica final de todos los "
        "nombres deberán ser confirmadas por las personas autoras antes del envío. A "
        "continuación se presentan las figuras maestras finales en inglés y sus leyendas "
        "completas, tal como se prepararon para el artículo.")
    set_run_font(note_run, size=Pt(BODY_FONT_SIZE))

    for paper_label, figure, legend in PAPER_FIGURES:
        document.add_page_break()
        add_heading(document, f"Artículo — {paper_label}", 2)
        width, height = picture_dimensions(
            figure, max_width=6.15, max_height=6.35)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.paragraph_format.space_after = Pt(4)
        image_paragraph.add_run().add_picture(
            str(figure), width=Inches(width), height=Inches(height))

        blocks = [re.sub(r"\s+", " ", block.strip())
                  for block in re.split(
                      r"\n\s*\n", legend.read_text(encoding="utf-8"))
                  if block.strip()]
        for position, block in enumerate(blocks):
            paragraph = add_caption(document, block)
            paragraph.paragraph_format.space_after = Pt(5)
            if position == 0 and paragraph.runs:
                paragraph.runs[0].bold = True


def add_results_and_discussion(document: Document) -> list[Path]:
    """Narrate each result before its sequentially numbered thesis figures."""
    records = thesis_panel_records()
    by_group: dict[int, dict[str, dict]] = {}
    for record in records:
        by_group.setdefault(record["group_index"], {})[record["letter"]] = record
    embedded: list[Path] = []

    def emit(group_index: int, letters: str, *, final: bool = False) -> None:
        selected = [by_group[group_index][letter] for letter in letters]
        for record in selected:
            add_panel_picture(
                document,
                record["panel"],
                f"Figura {record['number']}.",
                record["caption"],
            )
            embedded.append(record["panel"])
        if final:
            start = by_group[group_index][min(by_group[group_index])]["number"]
            end = by_group[group_index][max(by_group[group_index])]["number"]
            note = recode_panel_references(
                group_general_note(selected[-1]["legend"]), start)
            if note:
                add_caption(document, f"Nota de las Figuras {start}–{end}. {note}")

    document.add_page_break()
    add_heading(document, "IV. RESULTADOS", 1)

    add_heading(document, RESULT_SECTIONS[0][0], 2)
    add_body(document, RESULT_SECTIONS[0][1][0] +
             " Las trayectorias y el desenlace del día 6 se presentan en las Figuras 9 y 10.")
    emit(0, "AB")
    add_body(document, RESULT_SECTIONS[0][1][1] +
             " Las trayectorias y el desenlace de peso se presentan en las Figuras 11 y 12.")
    emit(0, "CD", final=True)

    add_heading(document, RESULT_SECTIONS[1][0], 2)
    add_body(
        document,
        "El protocolo temporal, la delimitación anatómica aceptada mediante HIL y la microscopía nativa "
        "que sustentan esta cuantificación se muestran en las Figuras 13–15.",
    )
    emit(1, "ABC")
    for text_value in RESULT_SECTIONS[1][1]:
        add_body(document, text_value)
    add_body(document, "Los resultados regionales por animal se resumen en la Figura 16.")
    emit(1, "D", final=True)

    add_heading(document, RESULT_SECTIONS[2][0], 2)
    add_body(
        document,
        "La microscopía trazable y las clasificaciones nucleares representativas de la cohorte "
        "NPY-GFP se presentan en las Figuras 17–19.",
    )
    emit(2, "ABC")
    add_body(document, RESULT_SECTIONS[2][1][0] +
             " Las dos variables de abundancia por animal se muestran en las Figuras 20 y 21.")
    emit(2, "DE")
    add_body(document, RESULT_SECTIONS[2][1][1] +
             " Los perfiles y mapas descriptivos se muestran en las Figuras 22 y 23.")
    emit(2, "FG", final=True)

    add_heading(document, RESULT_SECTIONS[3][0], 2)
    add_body(
        document,
        "La identidad de canales, la segmentación y la microscopía registrada usadas para el "
        "análisis POMC/c-FOS se documentan en las Figuras 24–29.",
    )
    emit(3, "ABCDEF")
    for text_value in RESULT_SECTIONS[3][1]:
        add_body(document, text_value)
    add_body(document, "Las variables c-FOS/DAPI y c-FOS/POMC por animal se presentan en las Figuras 30 y 31.")
    emit(3, "GH")
    add_body(
        document,
        "Los perfiles espaciales de c-FOS y de doble positividad c-FOS/POMC fueron descriptivos "
        "y no aportaron una señal confirmatoria; se presentan en las Figuras 32 y 33.",
    )
    emit(3, "IJ", final=True)

    add_heading(document, RESULT_SECTIONS[4][0], 2)
    add_body(
        document,
        "Las micrografías, máscaras HIL aceptadas, flujo de clasificación y esquemas de "
        "cuantificación glial se presentan en las Figuras 34–39.",
    )
    emit(4, "ABCDEF")
    add_body(document, RESULT_SECTIONS[4][1][0] +
             " La señal Iba1 regional por animal se muestra en la Figura 40.")
    emit(4, "G")
    add_body(document, RESULT_SECTIONS[4][1][1] +
             " La fracción morfométrica exploratoria se muestra en la Figura 41.")
    emit(4, "H")
    add_body(document, "La señal GFAP regional por animal, que tampoco difirió entre tratamientos, se muestra en la Figura 42.")
    emit(4, "I", final=True)

    add_heading(document, RESULT_SECTIONS[5][0], 2)
    add_body(document, RESULT_SECTIONS[5][1][0])
    add_body(document, "Los análisis masculinos se presentan en las Figuras 43–46.")
    emit(5, "ABCD", final=True)
    add_body(document, "Los análisis femeninos se presentan en las Figuras 47–50.")
    emit(6, "ABCD", final=True)
    add_body(document, RESULT_SECTIONS[5][1][1] +
             " El flujo H&E y sus análisis se presentan en las Figuras 51–56.")
    emit(7, "ABCDEF", final=True)
    add_body(document, RESULT_SECTIONS[5][1][2] +
             " Los cinco desenlaces del experimento de dos botellas se presentan en las Figuras 57–61.")
    emit(8, "ABCDE", final=True)
    add_body(document, RESULT_SECTIONS[5][1][3] +
             " La documentación y los análisis de esta cohorte se presentan en las Figuras 62–71.")
    emit(9, "ABCDEFGHIJ", final=True)

    if embedded != [record["panel"] for record in records]:
        raise RuntimeError("Los paneles españoles no se insertaron una sola vez y en orden canónico")

    document.add_page_break()
    add_heading(document, "V. DISCUSIÓN", 1)
    for title, paragraphs in DISCUSSION_SECTIONS:
        add_heading(document, title, 2)
        for text_value in paragraphs:
            add_body(document, text_value)

    document.add_page_break()
    add_heading(document, "VI. CONCLUSIONES", 1)
    for number, text_value in enumerate(CONCLUSIONS, 1):
        paragraph = add_body(document, f"{number}. {text_value}")
        paragraph.paragraph_format.first_line_indent = Inches(0)
    return embedded


def add_supplement(document: Document) -> None:
    document.add_page_break()
    add_heading(document, "ANEXO DE FIGURAS SUPLEMENTARIAS", 1)
    add_body(document, "Este anexo incorpora todos los subpaneles canónicos en español y sus leyendas completas. Los paneles proceden de la reconstrucción reproducible del repositorio y no de las figuras preliminares incrustadas en la versión anterior de la tesis.")
    for title, panels, legend, prefix in SUPPLEMENTARY_FIGURES:
        add_figure_set(document, title, panels, legend, prefix)


def write_audits(output: Path, panels: list[Path], reference_count: int,
                 cited_reference_count: int,
                 source_sha: str) -> tuple[Path, Path, Path]:
    audit_path = MANUSCRIPT_SOURCES / "TESIS_FINAL_NS4_sentence_audit_20260830.csv"
    manifest_path = MANUSCRIPT_SOURCES / "TESIS_FINAL_NS4_spanish_panel_manifest_20260830.csv"
    receipt_path = MANUSCRIPT_SOURCES / "TESIS_FINAL_NS4_completion_receipt_20260830.json"
    with audit_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["section", "item", "status", "basis"])
        for index in sorted(INTRO_REPLACEMENTS):
            writer.writerow(["Introducción", index, "revisado", "literatura y alcance corregidos"])
        for index in sorted(METHOD_REPLACEMENTS):
            writer.writerow(["Metodología", index, "revisado", "procedencia y análisis final"])
        for title, paragraphs in RESULT_SECTIONS:
            for item, _ in enumerate(paragraphs, 1):
                writer.writerow(["Resultados: " + title, item, "reemplazado", "tablas fuente canónicas"])
        for title, paragraphs in DISCUSSION_SECTIONS:
            for item, _ in enumerate(paragraphs, 1):
                writer.writerow(["Discusión: " + title, item, "revisado", "resultados definitivos y literatura"])
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["index", "relative_path", "sha256", "bytes"])
        for index, panel in enumerate(panels, 1):
            writer.writerow([index, panel.relative_to(ROOT), sha256(panel), panel.stat().st_size])
    receipt = {
        "status": "complete",
        "built_at": datetime.now(timezone.utc).isoformat(),
        "source": str(SOURCE),
        "source_sha256": source_sha,
        "output": str(output),
        "output_sha256": sha256(output),
        "spanish_panels": len(panels),
        "thesis_figure_numbering": {
            "preserved_and_method_figures": "1-8",
            "spanish_paper_panels": "9-71",
            "panel_letter_suffixes_in_thesis": False,
        },
        "paper_title": PAPER_TITLE,
        "paper_authors": PAPER_AUTHORS,
        "paper_appendix_master_figures": len(PAPER_FIGURES),
        "references": reference_count,
        "cited_references": cited_reference_count,
        "citation_style": "numeric_brackets_first_appearance",
        "thesis_format": {
            "font": "Times New Roman",
            "body_pt": BODY_FONT_SIZE,
            "figure_caption_pt": BODY_FONT_SIZE,
            "line_spacing": THESIS_LINE_SPACING,
            "title_pt_bold": TITLE_FONT_SIZE,
            "subtitle_pt_bold": SUBTITLE_FONT_SIZE,
        },
        "preserved_thesis_assets": ["esquemas hipotalámicos", "estructuras químicas", "diseño de dos botellas", "geles de genotipificación"],
        "replaced_preliminary_results": True,
        "raw_data_touched": False,
    }
    receipt_path.write_text(json.dumps(receipt, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return audit_path, manifest_path, receipt_path


def build(output: Path, page_map_path: Path | None = None) -> dict:
    global CITATIONS
    CITATIONS = CitationManager(load_references())
    panels = validate_inputs()
    source_sha = sha256(SOURCE)
    document = Document(SOURCE)
    original_paragraphs = list(document.paragraphs)
    if len(original_paragraphs) < 347:
        raise RuntimeError("La estructura del Word fuente no coincide con la plantilla preservada")

    style_document(document)
    page_map = load_page_map(page_map_path)
    fill_front_matter(document, original_paragraphs, page_map)
    apply_source_revisions(document, original_paragraphs)
    delete_from(document, 346)
    embedded = add_results_and_discussion(document)
    if embedded != panels:
        raise RuntimeError("El orden final de paneles no coincide con el manifiesto canónico")
    reference_records = CITATIONS.reference_records()
    cited_reference_count = len(CITATIONS.cited_keys)
    reference_count = add_references(document, reference_records)
    add_paper_appendix(document)
    if len(CITATIONS.cited_keys) != cited_reference_count:
        raise RuntimeError(
            "El anexo añadió citas después de construir la bibliografía")
    sanitize_fonts(document)
    add_update_fields_setting(document)

    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".building")
    document.core_properties.title = "Rol de la alulosa en la activación neuronal hipotalámica"
    document.core_properties.subject = "Tesis de Magíster completada con figuras canónicas en español"
    document.core_properties.comments = "Reconstruida el 30-08-2026; no modifica datos crudos."
    document.save(temporary)
    with zipfile.ZipFile(temporary) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"Miembro DOCX dañado: {bad_member}")
    check = Document(temporary)
    if len(check.inline_shapes) < 80:
        raise RuntimeError(f"Se esperaban al menos 80 imágenes; se hallaron {len(check.inline_shapes)}")
    temporary.replace(output)
    audit_path, manifest_path, receipt_path = write_audits(
        output, panels, reference_count, cited_reference_count, source_sha)
    return {
        "status": "complete",
        "output": str(output),
        "sha256": sha256(output),
        "bytes": output.stat().st_size,
        "paragraphs": len(check.paragraphs),
        "inline_shapes": len(check.inline_shapes),
        "spanish_panels": len(panels),
        "sequential_thesis_figures": "1-71",
        "paper_appendix_figures": len(PAPER_FIGURES),
        "references": reference_count,
        "cited_references": cited_reference_count,
        "citation_style": "numeric_brackets_first_appearance",
        "audit": str(audit_path),
        "panel_manifest": str(manifest_path),
        "receipt": str(receipt_path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--page-map", type=Path, default=PAGE_MAP_PATH)
    args = parser.parse_args()
    print(json.dumps(
        build(args.output, args.page_map), indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
